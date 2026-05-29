#!/usr/bin/env python3
from __future__ import annotations

import argparse
import asyncio
import bisect
import csv
import datetime as dt
import json
import logging
import math
import os
import queue
import re
import struct
import subprocess
import sys
import wave
import signal
import socket
import threading
import time
import tkinter as tk
from collections import Counter, deque
from dataclasses import dataclass, field
from shutil import which
from io import BytesIO
from pathlib import Path
from tkinter import ttk
from typing import Any, Optional

from aiohttp import web
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
import websockets


CR = "\r"
CR_BYTES = b"\r"
DEFAULT_CROSSMGR_PORT = 8765
DEFAULT_JCHIP_PORT = 53136
DEFAULT_WEB_PORT = 8865
RECENT_WINDOW_SECONDS = 10.0
DISTANCE_TO_FINISH_METERS = 50.0
MAX_PASSINGS = 500
GROUP_GAP_SECONDS = 2.0
MAX_GROUPS = 200
GROUP_MAX_AGE_SECONDS = 30.0
LOCAL_READ_DEDUP_SECONDS = 8.0
EXPORT_READ_DEDUP_SECONDS = LOCAL_READ_DEDUP_SECONDS
RE_TIME = re.compile(r"^\d\d:\d\d:\d\d\.\d+")
RE_SPEED = re.compile(r"([0-9]+(?:\.[0-9]+)?)\s*(km/h|kph|mph)?", re.I)
STATUS_TIMEOUTS = {
    "rfid": 15.0,
    "announcer": 15.0,
    "lapcounter": 15.0,
}


def ordinal(n: int) -> str:
    if 10 <= (n % 100) <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def _format_clock_precise(seconds: Optional[float], decimals: int) -> str:
    if seconds is None:
        return ""
    sign = "-" if seconds < 0 else ""
    seconds = abs(seconds)
    scale = 10 ** decimals
    total_units = int(round(seconds * scale))
    total_seconds, frac = divmod(total_units, scale)
    h, rem = divmod(total_seconds, 3600)
    m, s = divmod(rem, 60)
    frac_text = f".{frac:0{decimals}d}" if decimals else ""
    if h:
        return f"{sign}{h}:{m:02d}:{s:02d}{frac_text}"
    return f"{sign}{m}:{s:02d}{frac_text}"


def format_clock(seconds: Optional[float]) -> str:
    return _format_clock_precise(seconds, 2)


def format_clock_ms(seconds: Optional[float]) -> str:
    return _format_clock_precise(seconds, 3)


def format_eta(seconds: Optional[float]) -> str:
    if seconds is None:
        return ""
    return format_clock(seconds)


def format_elapsed_hms(seconds: Optional[float]) -> str:
    if seconds is None:
        return ""
    total = max(0, int(seconds))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


def parse_clock_to_days(value: str) -> Optional[float]:
    raw = (value or "").strip()
    if not raw:
        return None
    sign = -1 if raw.startswith("-") else 1
    clean = raw[1:] if sign < 0 else raw
    parts = clean.split(":")
    if len(parts) == 3:
        hours = int(parts[0])
        minutes = int(parts[1])
        seconds = float(parts[2])
    elif len(parts) == 2:
        hours = 0
        minutes = int(parts[0])
        seconds = float(parts[1])
    else:
        raise ValueError(f"unsupported time format: {value!r}")
    total_seconds = sign * (hours * 3600 + minutes * 60 + seconds)
    return total_seconds / 86400.0


def set_time_cell(ws: Any, ref: str, value: str, decimals: int = 3) -> None:
    parsed = parse_clock_to_days(value) if (value or "").strip() else None
    cell = ws[ref]
    cell.value = parsed
    if parsed is not None:
        cell.number_format = "[h]:mm:ss" + (("." + ("0" * decimals)) if decimals else "")


def parse_speed_mps(value: Any) -> Optional[float]:
    if not value:
        return None
    match = RE_SPEED.search(str(value))
    if not match:
        return None
    speed = float(match.group(1))
    unit = (match.group(2) or "km/h").lower()
    if unit in {"km/h", "kph"}:
        return speed / 3.6
    if unit == "mph":
        return speed * 0.44704
    return None


def as_utc_iso(value: dt.datetime) -> str:
    if value.tzinfo is None:
        value = value.astimezone()
    return value.isoformat()


def parse_host_port(value: str, default_port: int) -> tuple[str, int]:
    if ":" not in value:
        return value, default_port
    host, port = value.rsplit(":", 1)
    return host, int(port)


def build_tone_pcm(frequency: float = 880.0, duration: float = 0.18, sample_rate: int = 44100, amplitude: int = 32767 // 3) -> tuple[bytes, int]:
    total_samples = int(sample_rate * duration)
    pcm = bytearray()
    for i in range(total_samples):
        sample = int(amplitude * math.sin(2.0 * math.pi * frequency * (i / sample_rate)))
        pcm.extend(struct.pack("<h", sample))
    return bytes(pcm), sample_rate


def build_wave_bytes(pcm_bytes: bytes, sample_rate: int) -> bytes:
    buffer = BytesIO()
    with wave.open(buffer, "wb") as wav_file:
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(sample_rate)
        wav_file.writeframes(pcm_bytes)
    return buffer.getvalue()


def pick_audio_player() -> Optional[list[str]]:
    candidates = [
        ("paplay", ["paplay"]),
        ("aplay", ["aplay", "-q"]),
        ("ffplay", ["ffplay", "-nodisp", "-autoexit", "-loglevel", "quiet", "-i", "-"]),
        ("play", ["play", "-q", "-t", "wav", "-"]),
    ]
    for name, cmd in candidates:
        if which(name):
            return cmd
    return None


@dataclass(slots=True)
class RiderInfo:
    bib: str
    first_name: str = ""
    last_name: str = ""
    team: str = ""
    category: str = ""
    tag: str = ""

    @property
    def full_name(self) -> str:
        return " ".join(p for p in [self.first_name, self.last_name] if p).strip() or self.bib


@dataclass(slots=True)
class CategoryConfig:
    name: str
    numbers: str = ""
    race_laps: Optional[int] = None
    race_distance: Optional[float] = None
    start_offset: str = ""
    gender: str = ""
    category_type: str = ""

    def merge_live(self, detail: dict[str, Any]) -> "CategoryConfig":
        return CategoryConfig(
            name=str(detail.get("name") or self.name or "").strip(),
            numbers=self.numbers,
            race_laps=int(detail["laps"]) if detail.get("laps") not in (None, "") else self.race_laps,
            race_distance=float(detail["raceDistance"]) if detail.get("raceDistance") not in (None, "") else self.race_distance,
            start_offset=str(detail.get("startOffset") if detail.get("startOffset") not in (None, "") else self.start_offset),
            gender=str(detail.get("gender") or self.gender or "").strip(),
            category_type=str(detail.get("catType") or self.category_type or "").strip(),
        )


@dataclass(slots=True)
class PassingRecord:
    bib: str
    tag: str
    seen_at: dt.datetime
    received_at: dt.datetime
    race_time: Optional[float]
    note: str
    gap: str
    lap: Optional[int]
    name: str
    category: str
    team: str
    eta_seconds: Optional[float]
    category_rank: Optional[int]
    reader: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {
            "bib": self.bib,
            "tag": self.tag,
            "seen_at": as_utc_iso(self.seen_at),
            "received_at": as_utc_iso(self.received_at),
            "time": format_clock(self.race_time),
            "note": self.note,
            "gap": self.gap,
            "lap": self.lap,
            "name": self.name,
            "category": self.category,
            "team": self.team,
            "eta": format_eta(self.eta_seconds),
            "category_rank": self.category_rank,
            "reader": self.reader,
        }


@dataclass(slots=True)
class SubsystemStatus:
    name: str
    endpoint: str = ""
    state: str = "waiting"
    updated_at: Optional[dt.datetime] = None
    detail: str = ""

    def snapshot(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "endpoint": self.endpoint,
            "state": self.state,
            "updated_at": as_utc_iso(self.updated_at) if self.updated_at else "",
            "detail": self.detail,
        }


@dataclass(slots=True)
class IPHealthStatus:
    address: str
    last_ok_at: Optional[dt.datetime] = None
    last_probe_at: Optional[dt.datetime] = None
    last_rtt_ms: Optional[float] = None
    last_probe_ok: Optional[bool] = None

    def snapshot(self) -> dict[str, Any]:
        now = dt.datetime.now().astimezone()
        probe_age = None
        if self.last_probe_at is not None:
            probe_age = (now - self.last_probe_at).total_seconds()
        if self.last_probe_ok is True and probe_age is not None and probe_age <= 2.5:
            state = "ok"
        elif probe_age is not None and probe_age <= 5.0:
            state = "missed"
        else:
            state = "down"
        return {
            "address": self.address,
            "state": state,
            "last_ok_at": as_utc_iso(self.last_ok_at) if self.last_ok_at else "",
            "last_probe_at": as_utc_iso(self.last_probe_at) if self.last_probe_at else "",
            "last_rtt_ms": self.last_rtt_ms,
            "last_probe_ok": self.last_probe_ok,
        }


@dataclass(slots=True)
class RiderRuntimeState:
    bib: str
    predicted_lap: Optional[int] = None
    last_local_read_race_time: Optional[float] = None
    last_local_read_seen_at: Optional[dt.datetime] = None
    visible_in_results: bool = True
    category_hint: str = ""


GROUP_TREE_HEIGHT = 14


class RaceConfig:
    def __init__(self) -> None:
        self.event_name = ""
        self.event_date = ""
        self.timezone = ""
        self.time_trial = False
        self.categories: dict[str, CategoryConfig] = {}
        self.riders_by_tag: dict[str, RiderInfo] = {}
        self.riders_by_bib: dict[str, RiderInfo] = {}

    @classmethod
    def load(cls, xlsx_path: Path) -> "RaceConfig":
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        cfg = cls()

        ws = wb["Registration"]
        headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(headers, row))
            bib = str(rec.get("Bib#") or "").strip()
            tag = str(rec.get("Tag") or "").strip().upper()
            if not bib:
                continue
            rider = RiderInfo(
                bib=bib,
                first_name=str(rec.get("FirstName") or "").strip(),
                last_name=str(rec.get("LastName") or "").strip(),
                team=str(rec.get("Team") or "").strip(),
                category=str(rec.get("Category") or "").strip(),
                tag=tag,
            )
            cfg.riders_by_bib[bib] = rider
            if tag:
                cfg.riders_by_tag[tag] = rider

        ws = wb["--CrossMgr-Categories"]
        headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(headers, row))
            name = str(rec.get("Name") or "").strip()
            if not name:
                continue
            cfg.categories[name] = CategoryConfig(
                name=name,
                numbers=str(rec.get("Numbers") or "").strip(),
                race_laps=int(rec["Race Laps"]) if rec.get("Race Laps") not in (None, "") else None,
                race_distance=float(rec["Race Distance"]) if rec.get("Race Distance") not in (None, "") else None,
                start_offset=str(rec.get("Start Offset") or "").strip(),
                gender=str(rec.get("Gender") or "").strip(),
                category_type=str(rec.get("Category Type") or "").strip(),
            )

        ws = wb["--CrossMgr-Properties"]
        headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)
        if row:
            rec = dict(zip(headers, row))
            cfg.event_name = str(rec.get("Event Name") or "").strip()
            cfg.event_date = str(rec.get("Event Date") or "").strip()
            cfg.timezone = str(rec.get("TimeZone") or "").strip()
            cfg.time_trial = bool(rec.get("Time Trial"))

        return cfg


class SharedState:
    def __init__(self, config: RaceConfig, show_unknown_tags: bool = False) -> None:
        self.config = config
        self.show_unknown_tags = show_unknown_tags
        self.lock = threading.RLock()
        self.passings: deque[PassingRecord] = deque(maxlen=MAX_PASSINGS)
        self.results_by_bib: dict[str, dict[str, Any]] = {}
        self.rider_runtime: dict[str, RiderRuntimeState] = {}
        self.category_details: dict[str, dict[str, Any]] = {}
        self.latest_lap_refresh: dict[str, Any] = {}
        self.reference: dict[str, Any] = {}
        self._last_version_count: Optional[int] = None
        self._last_cur_race_time: Optional[float] = None
        self.reader_status: dict[str, str] = {}
        self.reader_clock_diff: dict[str, float] = {}
        self.race_clock_wall: Optional[float] = None
        self.race_clock_value: Optional[float] = None
        self.last_announcer_update: Optional[dt.datetime] = None
        self.last_lapcounter_update: Optional[dt.datetime] = None
        self.started_at = dt.datetime.now().astimezone()
        self.subsystems: dict[str, SubsystemStatus] = {
            "rfid": SubsystemStatus(name="RFID", endpoint="JChip :53136"),
            "announcer": SubsystemStatus(name="Announcer"),
            "lapcounter": SubsystemStatus(name="LapCounter"),
        }
        self.ip_health: dict[str, IPHealthStatus] = {}
        self.warning_enabled = True
        self.tt_export_path: Optional[Path] = None
        self.tt_xlsx_path: Optional[Path] = None
        self.tt_exported_keys: set[str] = set()
        self.tt_exported_read_times: dict[str, list[float]] = {}
        self.tt_diag_logged: set[str] = set()
        self.road_export_path: Optional[Path] = None
        self.road_xlsx_path: Optional[Path] = None
        self.road_pending_passings: list[PassingRecord] = []
        self.road_exported_keys: set[str] = set()
        self.road_exported_read_times: dict[str, list[float]] = {}
        self.road_exported_recorded_keys: set[str] = set()

    def effective_category_config(self, category: str) -> CategoryConfig:
        base = self.config.categories.get(category, CategoryConfig(name=category))
        detail = self._get_category_detail(category)
        if detail:
            return base.merge_live(detail)
        return base

    def current_race_time(self) -> Optional[float]:
        with self.lock:
            if self.race_clock_value is None:
                return None
            if self.race_clock_wall is None:
                return self.race_clock_value
            return max(0.0, self.race_clock_value + (time.time() - self.race_clock_wall))

    def is_time_trial(self) -> bool:
        with self.lock:
            return bool(self.config.time_trial or self.reference.get("isTimeTrial"))

    def tt_start_offset(self, bib: str) -> Optional[float]:
        result = self.results_by_bib.get(str(bib), {})
        if result.get("startTime") not in (None, ""):
            try:
                return float(result.get("startTime"))
            except (TypeError, ValueError):
                pass
        rider = self.config.riders_by_bib.get(str(bib))
        rider_start = getattr(rider, "start_time", None)
        return float(rider_start) if rider_start is not None else None

    def tt_export_csv_path(self) -> Optional[Path]:
        return self.tt_export_path

    def road_export_csv_path(self) -> Optional[Path]:
        return self.road_export_path

    def set_subsystem_status(self, key: str, state: str, endpoint: Optional[str] = None, detail: str = "") -> None:
        with self.lock:
            status = self.subsystems[key]
            if endpoint:
                status.endpoint = endpoint
            status.state = state
            status.updated_at = dt.datetime.now().astimezone()
            status.detail = detail

    def _effective_subsystem_state(self, key: str) -> SubsystemStatus:
        with self.lock:
            status = self.subsystems[key]
            effective = SubsystemStatus(
                name=status.name,
                endpoint=status.endpoint,
                state=status.state,
                updated_at=status.updated_at,
                detail=status.detail,
            )
            timeout = STATUS_TIMEOUTS.get(key)
            if timeout and effective.updated_at and effective.state == "connected":
                age = (dt.datetime.now().astimezone() - effective.updated_at).total_seconds()
                if age > timeout:
                    effective.state = "timedout"
            return effective

    def set_reader_status(self, reader: str, status: str) -> None:
        with self.lock:
            self.reader_status[reader] = status
        mapped = "connected" if status == "connected" else "timedout"
        self.set_subsystem_status("rfid", mapped, detail=reader)

    def set_reader_clock_diff(self, reader: str, delta_seconds: float) -> None:
        with self.lock:
            self.reader_clock_diff[reader] = delta_seconds

    def _reset_for_new_race_locked(self) -> None:
        self.passings.clear()
        self.results_by_bib.clear()
        self.category_details.clear()
        self.latest_lap_refresh = {}
        self.rider_runtime.clear()
        self.last_announcer_update = None
        self.last_lapcounter_update = None
        self.tt_exported_keys.clear()
        self.tt_exported_read_times.clear()
        self.tt_diag_logged.clear()
        self.road_pending_passings.clear()
        self.road_exported_keys.clear()
        self.road_exported_read_times.clear()
        self.road_exported_recorded_keys.clear()

    def reset_for_reader_reconnect(self, reader: str = "") -> None:
        with self.lock:
            self.passings.clear()
            self.rider_runtime.clear()
            self.road_pending_passings.clear()
        logging.getLogger("lapsrv.state").warning(
            "reader reconnect reset local RFID state%s",
            f": {reader}" if reader else "",
        )

    def update_reference(self, ref: dict[str, Any]) -> None:
        with self.lock:
            new_version_count = None
            if (ref or {}).get("versionCount") is not None:
                try:
                    new_version_count = int((ref or {}).get("versionCount"))
                except (TypeError, ValueError):
                    new_version_count = None
            new_cur_race_time = None
            if (ref or {}).get("curRaceTime") is not None:
                try:
                    new_cur_race_time = float((ref or {}).get("curRaceTime"))
                except (TypeError, ValueError):
                    new_cur_race_time = None

            should_reset = False
            if self._last_version_count is not None and new_version_count is not None and new_version_count < self._last_version_count:
                should_reset = True
            elif (
                self._last_cur_race_time is not None
                and new_cur_race_time is not None
                and self._last_cur_race_time > 30.0
                and new_cur_race_time < 5.0
                and new_cur_race_time + 10.0 < self._last_cur_race_time
            ):
                should_reset = True

            if should_reset:
                logging.getLogger("lapsrv.state").info(
                    "detected new race, resetting volatile state: version %s->%s race_time %.2f->%.2f",
                    self._last_version_count,
                    new_version_count,
                    self._last_cur_race_time if self._last_cur_race_time is not None else -1.0,
                    new_cur_race_time if new_cur_race_time is not None else -1.0,
                )
                self._reset_for_new_race_locked()
                self.reference = {}

            self.reference.update(ref or {})
            cur_race_time = self.reference.get("curRaceTime")
            if cur_race_time is not None:
                try:
                    new_clock_value = float(cur_race_time)
                    if self.race_clock_value is None or abs(new_clock_value - self.race_clock_value) > 1e-6:
                        self.race_clock_value = new_clock_value
                        self.race_clock_wall = time.time()
                except (TypeError, ValueError):
                    pass
            self._last_version_count = new_version_count if new_version_count is not None else self._last_version_count
            self._last_cur_race_time = new_cur_race_time if new_cur_race_time is not None else self._last_cur_race_time

    def update_announcer(self, msg: dict[str, Any]) -> None:
        cmd = msg.get("cmd")
        ref = msg.get("reference") or {}
        self.update_reference(ref)
        with self.lock:
            if cmd == "baseline":
                self.results_by_bib = {str(k): v for k, v in (msg.get("info") or {}).items()}
                self.category_details = {str(k): v for k, v in (msg.get("categoryDetails") or {}).items()}
            elif cmd == "ram":
                self._apply_ram(self.results_by_bib, msg.get("infoRAM") or {})
                self._apply_ram(self.category_details, msg.get("categoryRAM") or {})
            self._refresh_rider_runtime_from_announcer()
            self.last_announcer_update = dt.datetime.now().astimezone()
            self.subsystems["announcer"].updated_at = self.last_announcer_update
        logging.getLogger("lapsrv.crossmgr").debug(
            "announcer %s: info=%d category_details=%d ref_keys=%s",
            cmd,
            len(self.results_by_bib),
            len(self.category_details),
            sorted((ref or {}).keys()),
        )

    def update_lapcounter(self, msg: dict[str, Any]) -> None:
        self.update_reference(msg)
        with self.lock:
            self.latest_lap_refresh = msg
            self.last_lapcounter_update = dt.datetime.now().astimezone()
            self.subsystems["lapcounter"].updated_at = self.last_lapcounter_update
        logging.getLogger("lapsrv.crossmgr").debug(
            "lapcounter refresh: labels=%s curRaceTime=%s",
            msg.get("labels"),
            msg.get("curRaceTime"),
        )

    @staticmethod
    def _apply_ram(dest: dict[str, Any], ram: dict[str, Any]) -> None:
        for key, value in (ram.get("a") or {}).items():
            dest[str(key)] = value
        for key, value in (ram.get("m") or {}).items():
            dest[str(key)] = value
        for key in (ram.get("r") or []):
            dest.pop(str(key), None)

    def _get_category_detail(self, category: str) -> Optional[dict[str, Any]]:
        detail = self.category_details.get(category)
        if detail:
            return detail
        for candidate in self.category_details.values():
            if isinstance(candidate, dict) and candidate.get("name") == category:
                return candidate
        return None

    def _sorted_category_details(self) -> list[dict[str, Any]]:
        details = [v for v in self.category_details.values() if isinstance(v, dict)]
        details.sort(key=lambda d: (int(d.get("iSort", 999999)), str(d.get("name") or "")))
        return details

    def _lapcounter_label_map(self) -> dict[str, list[Any]]:
        labels = self.latest_lap_refresh.get("labels") or []
        if not labels:
            return {}
        categories = self._sorted_category_details()
        if len(categories) == len(labels) + 1:
            categories = categories[1:]
        elif len(categories) > len(labels):
            categories = categories[: len(labels)]
        elif len(labels) > len(categories):
            labels = labels[: len(categories)]
        mapping: dict[str, list[Any]] = {}
        for detail, label_row in zip(categories, labels):
            name = str(detail.get("name") or "").strip()
            if name:
                mapping[name] = label_row
        return mapping

    def _lapcounter_display_for_category(self, category: str) -> Optional[dict[str, Any]]:
        label_row = self._lapcounter_label_map().get(category)
        if not label_row:
            return None
        label = str(label_row[0]) if len(label_row) >= 1 and label_row[0] is not None else ""
        flash = bool(label_row[1]) if len(label_row) >= 2 else False
        lap_start = float(label_row[2]) if len(label_row) >= 3 and label_row[2] is not None else None
        return {"label": label, "flash": flash, "lap_start": lap_start}

    def _rider_state(self, bib: str) -> RiderRuntimeState:
        bib = str(bib)
        state = self.rider_runtime.get(bib)
        if state is None:
            state = RiderRuntimeState(bib=bib)
            self.rider_runtime[bib] = state
        return state

    def _refresh_rider_runtime_from_announcer(self) -> None:
        visible_bibs: set[str] = set()
        for candidate in self.category_details.values():
            if not isinstance(candidate, dict):
                continue
            name = str(candidate.get("name") or "").strip()
            if not name or name.lower() == "all":
                continue
            for bib in (candidate.get("pos") or []):
                bib = str(bib)
                visible_bibs.add(bib)
                state = self._rider_state(bib)
                state.visible_in_results = True
                state.category_hint = name

        if visible_bibs:
            all_bibs = set(self.rider_runtime) | set(self.results_by_bib) | set(self.config.riders_by_bib)
            for bib in all_bibs:
                state = self._rider_state(bib)
                if bib not in visible_bibs:
                    state.visible_in_results = False

        current_race_time = self.current_race_time()
        if current_race_time is None:
            return
        for bib in set(self.results_by_bib) | visible_bibs:
            progress = self._authoritative_progress_for_bib(bib, current_race_time)
            auth_lap = progress.get("recorded_lap")
            if auth_lap is None:
                continue
            state = self._rider_state(bib)
            if state.predicted_lap is None or auth_lap >= state.predicted_lap:
                state.predicted_lap = int(auth_lap)

    def _category_name_for_bib(self, bib: str) -> Optional[str]:
        bib = str(bib)
        state = self.rider_runtime.get(bib)
        if state and state.category_hint and state.category_hint.lower() != "all":
            return state.category_hint
        for candidate in self.category_details.values():
            if not isinstance(candidate, dict):
                continue
            name = str(candidate.get("name") or "").strip()
            if not name or name.lower() == "all":
                continue
            pos = [str(v) for v in (candidate.get("pos") or [])]
            if bib in pos:
                return name
        return None

    def _effective_category_name(self, rider: RiderInfo, result: dict[str, Any]) -> str:
        result_category = str(result.get("category") or "").strip()
        if result_category.lower() == "all":
            result_category = ""
        result_category_name = str(result.get("categoryName") or "").strip()
        if result_category_name.lower() == "all":
            result_category_name = ""
        return (
            self._category_name_for_bib(rider.bib)
            or result_category_name
            or result_category
            or rider.category
            or "Unknown"
        )

    def _race_times(self, result: dict[str, Any]) -> list[float]:
        return [float(v) for v in (result.get("raceTimes") or []) if v is not None]

    def _interp_flags(self, result: dict[str, Any]) -> list[bool]:
        return [bool(v) for v in (result.get("interp") or [])]

    def _authoritative_progress_for_bib(self, bib: str, current_race_time: Optional[float]) -> dict[str, Optional[float]]:
        result = self.results_by_bib.get(str(bib))
        if not result:
            return {"recorded_lap": None, "recorded_time": None, "expected_lap": None, "expected_time": None}
        return self._result_progress(result, current_race_time)

    def _announcer_recorded_map(self, current_race_time: Optional[float]) -> dict[str, dict[str, Any]]:
        if current_race_time is None:
            return {}
        recorded: dict[str, dict[str, Any]] = {}
        is_time_trial = bool(self.reference.get("isTimeTrial") or self.config.time_trial)
        for bib, result in self.results_by_bib.items():
            if not result:
                continue
            status = str(result.get("status") or "").strip()
            race_times = self._race_times(result)
            interp = self._interp_flags(result)
            if is_time_trial:
                if status not in {"Finisher", "NP"} or not race_times:
                    continue
                last_index = len(race_times) - 1
                if last_index < 0:
                    continue
                if last_index < len(interp) and interp[last_index]:
                    continue
                offset = float(result.get("startTime") or 0.0)
                recorded[str(bib)] = {"lap": int(last_index), "t": float(race_times[last_index] + offset)}
                continue
            if status != "Finisher" or len(race_times) < 2:
                continue
            progress = self._result_progress(result, current_race_time)
            recorded_lap = progress.get("recorded_lap")
            recorded_time = progress.get("recorded_time")
            if recorded_lap is None or recorded_time is None or recorded_lap < 1:
                continue
            recorded[str(bib)] = {"lap": int(recorded_lap), "t": float(recorded_time)}
        return recorded

    def _road_recorded_events_by_bib(self) -> dict[str, list[dict[str, Any]]]:
        recorded: dict[str, list[dict[str, Any]]] = {}
        for bib, result in self.results_by_bib.items():
            if not result:
                continue
            race_times = self._race_times(result)
            interp = self._interp_flags(result)
            if len(race_times) < 2:
                continue
            events: list[dict[str, Any]] = []
            for lap in range(1, len(race_times)):
                if lap < len(interp) and interp[lap]:
                    continue
                events.append({"lap": lap, "t": float(race_times[lap])})
            if events:
                recorded[str(bib)] = sorted(events, key=lambda item: float(item["t"]))
        return recorded

    def _result_progress(self, result: dict[str, Any], current_race_time: Optional[float]) -> dict[str, Optional[float]]:
        if current_race_time is None or not result:
            return {"recorded_lap": None, "recorded_time": None, "expected_lap": None, "expected_time": None}
        race_times = self._race_times(result)
        if not race_times:
            return {"recorded_lap": 0, "recorded_time": None, "expected_lap": None, "expected_time": None}
        is_time_trial = bool(self.reference.get("isTimeTrial"))
        offset = float(result.get("startTime") or 0.0) if is_time_trial else 0.0
        target = current_race_time - offset
        i = bisect.bisect_left(race_times, target)
        interp = self._interp_flags(result)

        expected_lap: Optional[int] = None
        expected_time: Optional[float] = None
        lap = i
        if lap > 1 and lap - 1 < len(interp) and interp[lap - 1]:
            lap -= 1
        if lap < len(race_times) and lap < len(interp) and interp[lap]:
            expected_lap = lap
            expected_time = race_times[lap] + offset

        lap = i - 1
        while lap > 0 and lap < len(interp) and interp[lap]:
            lap -= 1
        recorded_lap = max(lap, 0)
        recorded_time = None
        if recorded_lap < len(race_times) and (recorded_lap == 0 or recorded_lap >= len(interp) or not interp[recorded_lap]):
            recorded_time = race_times[recorded_lap] + offset

        return {
            "recorded_lap": recorded_lap,
            "recorded_time": recorded_time,
            "expected_lap": expected_lap,
            "expected_time": expected_time,
        }

    def _lap_for_rider(self, bib: str, current_race_time: Optional[float]) -> Optional[int]:
        bib = str(bib)
        state = self._rider_state(bib)
        progress = self._authoritative_progress_for_bib(bib, current_race_time)
        lap = progress.get("recorded_lap")
        if state.predicted_lap is not None and (lap is None or state.predicted_lap > lap):
            return int(state.predicted_lap)
        return int(lap) if lap is not None else state.predicted_lap

    def _rank_for_rider(self, rider: RiderInfo) -> Optional[int]:
        result = self.results_by_bib.get(rider.bib, {})
        category = self._effective_category_name(rider, result)
        detail = self._get_category_detail(category)
        if not detail:
            return None
        pos = [str(v) for v in (detail.get("pos") or [])]
        try:
            return pos.index(rider.bib) + 1
        except ValueError:
            return None

    def _category_bibs(self, category: str) -> list[str]:
        bibs: list[str] = []
        seen: set[str] = set()
        detail = self._get_category_detail(category)
        if detail:
            for bib in (detail.get("pos") or []):
                bib_s = str(bib)
                if bib_s and bib_s not in seen:
                    seen.add(bib_s)
                    bibs.append(bib_s)
        for bib, rider in self.config.riders_by_bib.items():
            result = self.results_by_bib.get(bib, {})
            effective = self._effective_category_name(rider, result)
            if effective == category and bib not in seen:
                seen.add(bib)
                bibs.append(bib)
        return bibs

    def _rank_sort_key_for_bib(self, bib: str, category: str, current_race_time: Optional[float]) -> tuple[int, float, int, str]:
        rider = self.config.riders_by_bib.get(str(bib))
        if not rider:
            return (999999, float("inf"), 999999, str(bib))
        lap = self._lap_for_rider(rider.bib, current_race_time)
        if lap is None:
            lap = -1
        state = self._rider_state(rider.bib)
        lap_entry_time: Optional[float] = state.last_local_read_race_time
        progress = self._authoritative_progress_for_bib(rider.bib, current_race_time)
        recorded_time = progress.get("recorded_time")
        recorded_lap = progress.get("recorded_lap")
        if lap_entry_time is None and recorded_time is not None and recorded_lap is not None:
            lap_entry_time = float(recorded_time)
        fallback_rank = self._rank_for_rider(rider) or 999999
        return (-int(lap), float(lap_entry_time) if lap_entry_time is not None else float("inf"), fallback_rank, rider.bib)

    def _display_rank_for_bib(self, bib: str, category: str, current_race_time: Optional[float]) -> Optional[int]:
        bibs = self._category_bibs(category)
        if not bibs:
            return None
        ordered = sorted(bibs, key=lambda b: self._rank_sort_key_for_bib(b, category, current_race_time))
        try:
            return ordered.index(str(bib)) + 1
        except ValueError:
            return None

    def _leader_lap(self, rider: RiderInfo, current_race_time: Optional[float]) -> Optional[int]:
        result = self.results_by_bib.get(rider.bib, {})
        category = self._effective_category_name(rider, result)
        detail = self._get_category_detail(category)
        if not detail or current_race_time is None:
            return None
        bibs = self._category_bibs(category)
        if not bibs:
            return None
        leader_bib = min(bibs, key=lambda b: self._rank_sort_key_for_bib(b, category, current_race_time))
        return self._lap_for_rider(leader_bib, current_race_time)

    def _note_for_rider(self, rider: RiderInfo) -> tuple[str, Optional[int]]:
        rank = self._display_rank_for_bib(rider.bib, rider.category or self._effective_category_name(rider, self.results_by_bib.get(rider.bib, {})), self.current_race_time())
        if rank == 1:
            return "1st", rank
        if rank:
            return ordinal(rank), rank
        return "", None

    def _rider_rank_and_lap_deficit(self, bib: str, current_race_time: Optional[float]) -> tuple[Optional[int], int]:
        rider = self.config.riders_by_bib.get(str(bib))
        if not rider:
            return None, 0
        category = self._effective_category_name(rider, self.results_by_bib.get(rider.bib, {}))
        rank = self._display_rank_for_bib(rider.bib, category, current_race_time)
        rider_lap = self._lap_for_rider(rider.bib, current_race_time)
        leader_lap = self._leader_lap(rider, current_race_time)
        lap_deficit = 0
        if rider_lap is not None and leader_lap is not None and leader_lap > rider_lap:
            lap_deficit = leader_lap - rider_lap
        return rank, lap_deficit

    def _group_note_for_category(
        self, group_passings: list[PassingRecord], category: str, current_race_time: Optional[float], lap_deficit_filter: Optional[int] = None
    ) -> tuple[str, bool, bool]:
        category_passings = [p for p in group_passings if p.category == category and p.bib not in {"?", ""}]
        ranks: list[int] = []
        deficits: list[int] = []
        best_rank: Optional[int] = None
        for passing in category_passings:
            rank, lap_deficit = self._rider_rank_and_lap_deficit(passing.bib, current_race_time)
            if not rank:
                continue
            if lap_deficit_filter is not None and lap_deficit != lap_deficit_filter:
                continue
            ranks.append(rank)
            deficits.append(lap_deficit)
            if best_rank is None or rank < best_rank:
                best_rank = rank
        if not ranks or best_rank is None:
            return "", False, bool(lap_deficit_filter and lap_deficit_filter > 0)
        worst_rank = max(ranks)
        if best_rank == worst_rank:
            note = "1st" if best_rank == 1 else ordinal(best_rank)
        else:
            start = "1st" if best_rank == 1 else ordinal(best_rank)
            note = f"{start}:{ordinal(worst_rank)}"
        all_lapped = bool(deficits) and min(deficits) > 0
        return note, best_rank == 1, all_lapped

    def _gap_for_rider(self, rider: RiderInfo, category: str, lap: Optional[int], current_race_time: Optional[float]) -> str:
        rank = self._rank_for_rider(rider)
        leader_lap = self._leader_lap(rider, current_race_time)
        if lap is not None and leader_lap is not None and leader_lap > lap:
            return f"-{leader_lap - lap} laps"
        if not rank or rank <= 1:
            return ""
        with self.lock:
            previous_same_cat = next(
                (p for p in reversed(self.passings) if p.category == category and p.bib != rider.bib),
                None,
            )
        if previous_same_cat and previous_same_cat.race_time is not None and current_race_time is not None:
            delta = abs((previous_same_cat.race_time or 0.0) - current_race_time)
            return format_clock(delta)
        detail = self._get_category_detail(category)
        if detail:
            pos = [str(v) for v in (detail.get("pos") or [])]
            gap_values = detail.get("gapValue") or []
            try:
                idx = pos.index(rider.bib)
                if idx < len(gap_values):
                    gap = gap_values[idx]
                    if isinstance(gap, (int, float)):
                        return format_clock(float(gap)) if float(gap) > 0 else str(gap)
                    if gap is not None:
                        return str(gap)
            except ValueError:
                pass
        return ""

    def _eta_seconds(self, rider: RiderInfo, category: str, result: dict[str, Any], lap: Optional[int]) -> Optional[float]:
        progress = self._authoritative_progress_for_bib(rider.bib, self.current_race_time())
        expected_time = progress.get("expected_time")
        current_race_time = self.current_race_time()
        if expected_time is not None and current_race_time is not None:
            return expected_time - current_race_time
        speed = parse_speed_mps(result.get("speed"))
        if speed and speed > 0:
            return DISTANCE_TO_FINISH_METERS / speed
        race_times = self._race_times(result)
        if lap is None or lap <= 0:
            return None
        lap_distance_km = result.get("lapDistance")
        if lap_distance_km in (None, ""):
            detail = self._get_category_detail(category)
            if detail:
                lap_distance_km = detail.get("lapDistance")
        cat_cfg = self.effective_category_config(category)
        if (not cat_cfg.name or cat_cfg.name == category) and rider.category and rider.category != category:
            fallback_cfg = self.config.categories.get(rider.category)
            if fallback_cfg:
                cat_cfg = fallback_cfg
        if lap_distance_km in (None, "") and cat_cfg and cat_cfg.race_distance:
            lap_distance_km = cat_cfg.race_distance
        if lap_distance_km in (None, "") or lap >= len(race_times):
            return None
        lap_time = race_times[lap] - race_times[lap - 1]
        if lap_time <= 0:
            return None
        lap_distance_m = float(lap_distance_km) * 1000.0
        speed = lap_distance_m / lap_time
        if speed <= 0:
            return None
        return DISTANCE_TO_FINISH_METERS / speed

    def record_passing(self, tag: str, seen_at: dt.datetime, reader: str = "") -> Optional[PassingRecord]:
        tag = tag.strip().upper()
        received_at = dt.datetime.now().astimezone()
        self.set_subsystem_status("rfid", "connected", detail=reader or tag)
        with self.lock:
            rider = self.config.riders_by_tag.get(tag)
            current_race_time = self.current_race_time()
            race_time = None
            if current_race_time is not None:
                if self.race_clock_wall is not None:
                    race_time = max(0.0, self.race_clock_value + (seen_at.timestamp() - self.race_clock_wall))
                else:
                    race_time = current_race_time

            if rider:
                state = self._rider_state(rider.bib)
                if race_time is not None:
                    auth_progress = self._authoritative_progress_for_bib(rider.bib, race_time)
                    auth_lap = auth_progress.get("recorded_lap")
                    new_predicted_lap = auth_lap if auth_lap is not None else state.predicted_lap
                    is_new_local_read = (
                        state.last_local_read_race_time is None
                        or abs(race_time - state.last_local_read_race_time) >= LOCAL_READ_DEDUP_SECONDS
                    )
                    if is_new_local_read:
                        if auth_lap is not None:
                            new_predicted_lap = max(new_predicted_lap or 0, auth_lap + 1)
                        elif state.predicted_lap is not None:
                            new_predicted_lap = state.predicted_lap + 1
                    state.predicted_lap = int(new_predicted_lap) if new_predicted_lap is not None else state.predicted_lap
                    state.last_local_read_race_time = race_time
                    state.last_local_read_seen_at = seen_at
                result = self.results_by_bib.get(rider.bib, {})
                category = self._effective_category_name(rider, result)
                if category and category != "Unknown":
                    state.category_hint = category
                if not state.visible_in_results:
                    return None
                lap = self._lap_for_rider(rider.bib, race_time if race_time is not None else current_race_time)
                note, rank = self._note_for_rider(rider)
                gap = self._gap_for_rider(rider, category, lap, race_time if race_time is not None else current_race_time)
                eta = self._eta_seconds(rider, category, result, lap)
                record = PassingRecord(
                    bib=rider.bib,
                    tag=tag,
                    seen_at=seen_at,
                    received_at=received_at,
                    race_time=race_time,
                    note=note,
                    gap=gap,
                    lap=lap,
                    name=rider.full_name,
                    category=category,
                    team=rider.team,
                    eta_seconds=eta,
                    category_rank=rank,
                    reader=reader,
                )
            else:
                if not self.show_unknown_tags:
                    return None
                record = PassingRecord(
                    bib=tag[:12],
                    tag=tag,
                    seen_at=seen_at,
                    received_at=received_at,
                    race_time=race_time,
                    note="Unknown tag",
                    gap="",
                    lap=None,
                    name="Unknown rider",
                    category="Unknown",
                    team="",
                    eta_seconds=None,
                    category_rank=None,
                    reader=reader,
                )
            self.passings.append(record)
            if not self.is_time_trial() and rider and record.race_time is not None:
                duplicate_pending = any(
                    p.bib == record.bib
                    and p.race_time is not None
                    and abs(float(p.race_time) - float(record.race_time)) < LOCAL_READ_DEDUP_SECONDS
                    for p in reversed(self.road_pending_passings[-20:])
                )
                if not duplicate_pending:
                    self.road_pending_passings.append(record)
            return record

    def _category_laps_to_go(self, category: str, current_race_time: Optional[float]) -> Optional[int]:
        detail = self._get_category_detail(category)
        if not detail:
            return None
        cfg = self.effective_category_config(category)
        total_laps = detail.get("laps")
        if total_laps in (None, "", 0):
            total_laps = cfg.race_laps if cfg and cfg.race_laps is not None else None
        if total_laps is None:
            return None
        pos = [str(v) for v in (detail.get("pos") or [])]
        if not pos:
            return int(total_laps)
        leader_result = self.results_by_bib.get(pos[0], {})
        leader_progress = self._result_progress(leader_result, current_race_time)
        leader_lap = leader_progress.get("expected_lap")
        if leader_lap is None:
            leader_lap = leader_progress.get("recorded_lap")
        if leader_lap is None:
            return int(total_laps)
        return max(int(total_laps) - leader_lap, 0)

    def _category_display_summary(self, category: str, current_race_time: Optional[float]) -> dict[str, Any]:
        lapcounter_display = self._lapcounter_display_for_category(category)
        laps_to_go = self._category_laps_to_go(category, current_race_time)
        lap_label = None
        if lapcounter_display:
            lap_label = lapcounter_display["label"]
            try:
                laps_to_go = int(lap_label)
            except (TypeError, ValueError):
                pass
        return {
            "laps_to_go": laps_to_go,
            "lap_label": lap_label,
            "bell": (lap_label == "1") if lap_label is not None else laps_to_go == 1,
        }

    def _category_display_order(self) -> list[str]:
        raw_names: list[str] = []
        raw_names.extend(name for name in self.config.categories.keys() if name and name.lower() != "all")
        for detail in self._sorted_category_details():
            name = str(detail.get("name") or "").strip()
            if name and name.lower() != "all":
                raw_names.append(name)

        names: list[str] = []
        seen: set[str] = set()
        for name in raw_names:
            if name in seen:
                continue
            seen.add(name)
            names.append(name)

        gendered_prefixes = {name.split(" (", 1)[0] for name in names if " (" in name}
        ordered: list[str] = []
        for name in names:
            if " (" not in name and name in gendered_prefixes:
                continue
            ordered.append(name)
        return ordered

    def _category_header_status(self, category: str, current_race_time: Optional[float]) -> dict[str, Any]:
        summary = self._category_display_summary(category, current_race_time)
        lap_text = summary["lap_label"] if summary["lap_label"] not in (None, "") else summary["laps_to_go"]
        if summary["bell"]:
            return {"text": "bell", "is_bell": True}
        return {"text": "" if lap_text in (None, "") else str(lap_text), "is_bell": False}

    def _build_group_tables(self) -> tuple[dict[str, Any], dict[str, Any]]:
        with self.lock:
            passings = list(self.passings)
            current_race_time = self.current_race_time()
            category_order = self._category_display_order()
        header_status = {category: self._category_header_status(category, current_race_time) for category in category_order}
        announcer_recorded = self._announcer_recorded_map(current_race_time)
        empty = {"headers": category_order, "header_status": header_status, "rows": []}
        if not passings:
            return empty, empty

        groups: list[dict[str, Any]] = []
        current_group: Optional[dict[str, Any]] = None
        prev_seen_at: Optional[dt.datetime] = None
        for passing in passings:
            if current_group is None:
                current_group = {"passings": [passing], "start": passing.seen_at, "end": passing.seen_at}
                groups.append(current_group)
            else:
                gap = (passing.seen_at - prev_seen_at).total_seconds() if prev_seen_at else 0.0
                if gap >= GROUP_GAP_SECONDS:
                    current_group = {"passings": [passing], "start": passing.seen_at, "end": passing.seen_at}
                    groups.append(current_group)
                else:
                    current_group["passings"].append(passing)
                    current_group["end"] = passing.seen_at
            prev_seen_at = passing.seen_at
        cutoff_seen_at = dt.datetime.now().astimezone() - dt.timedelta(seconds=GROUP_MAX_AGE_SECONDS)
        groups = [group for group in groups if group.get("end") and group["end"] >= cutoff_seen_at]
        groups = groups[-MAX_GROUPS:]

        rows: list[dict[str, Any]] = []
        previous_group_time: Optional[float] = None

        def build_group_row(group_passings: list[PassingRecord], group_race_time: Optional[float], gap_seconds: Optional[float], is_past: bool) -> Optional[dict[str, Any]]:
            if not group_passings:
                return None
            per_category: dict[str, dict[str, Any]] = {}
            for passing in group_passings:
                category = passing.category
                if category not in category_order:
                    if any(name.startswith(f"{category} (") for name in category_order):
                        continue
                    category_order.append(category)
                _rank, lap_deficit = self._rider_rank_and_lap_deficit(passing.bib, current_race_time)
                info = per_category.setdefault(category, {"count": 0, "bibs": [], "deficit_counts": {}})
                info["count"] += 1
                info["deficit_counts"][lap_deficit] = int(info["deficit_counts"].get(lap_deficit, 0)) + 1
                if passing.bib not in info["bibs"]:
                    info["bibs"].append(passing.bib)
            cells: dict[str, dict[str, Any]] = {}
            for category, info in per_category.items():
                count = int(info["count"])
                note, is_lead, is_lapped = self._group_note_for_category(group_passings, category, current_race_time)
                cells[category] = {
                    "text": note,
                    "note": note,
                    "count": count,
                    "is_lead": is_lead,
                    "is_lapped": is_lapped,
                    "deficit_counts": dict(info.get("deficit_counts") or {}),
                }
            visible_cells = {category: cells[category] for category in category_order if category in cells}
            if not visible_cells:
                return None
            return {
                "elapsed_text": format_elapsed_hms(group_race_time),
                "group_gap": format_elapsed_hms(gap_seconds) if gap_seconds is not None else "",
                "total": sum(cell.get("count", 0) for cell in visible_cells.values()),
                "cells": visible_cells,
                "is_past": is_past,
            }

        for group in groups:
            group_passings: list[PassingRecord] = group["passings"]
            group_race_time = next((p.race_time for p in group_passings if p.race_time is not None), None)
            gap_seconds = None
            if group_race_time is not None and previous_group_time is not None:
                gap_seconds = group_race_time - previous_group_time
            if group_race_time is not None:
                previous_group_time = group_race_time

            active_passings: list[PassingRecord] = []
            past_passings: list[PassingRecord] = []
            for passing in group_passings:
                rec = announcer_recorded.get(str(passing.bib))
                if rec and (group_race_time is None or rec.get("t") >= group_race_time):
                    past_passings.append(passing)
                else:
                    active_passings.append(passing)

            recent_row = build_group_row(active_passings, group_race_time, gap_seconds, False)
            if recent_row is not None:
                rows.append(recent_row)
            past_row = build_group_row(past_passings, group_race_time, gap_seconds, True)
            if past_row is not None:
                rows.append(past_row)
        header_status = {category: self._category_header_status(category, current_race_time) for category in category_order}
        recent_rows = [dict(r) for r in rows if not r.get("is_past")]
        past_rows = [dict(r) for r in rows if r.get("is_past")]
        past_rows.reverse()
        for r in recent_rows:
            r.pop("is_past", None)
        for r in past_rows:
            r.pop("is_past", None)
        return (
            {"headers": category_order, "header_status": header_status, "rows": recent_rows},
            {"headers": category_order, "header_status": header_status, "rows": past_rows},
        )

    def _tt_row_from_passing(self, passing: PassingRecord, official_race_time: Optional[float]) -> dict[str, Any]:
        rider = self.config.riders_by_bib.get(str(passing.bib))
        start_time = self.tt_start_offset(str(passing.bib))
        race_time = passing.race_time
        elapsed_base = official_race_time if official_race_time is not None else race_time
        elapsed = None if elapsed_base is None or start_time is None else max(0.0, elapsed_base - start_time)
        display_race_time = official_race_time if official_race_time is not None else race_time
        return {
            "race_time": format_clock(display_race_time),
            "race_time_csv": format_clock_ms(display_race_time),
            "early_time": format_clock(race_time),
            "early_time_csv": format_clock_ms(race_time),
            "start_time": format_clock(start_time),
            "start_time_csv": format_clock_ms(start_time),
            "stop_time": format_clock(official_race_time),
            "stop_time_csv": format_clock_ms(official_race_time),
            "elapsed": format_clock(elapsed),
            "elapsed_csv": format_clock_ms(elapsed),
            "bib": str(passing.bib),
            "last_name": rider.last_name if rider else "",
            "first_name": rider.first_name if rider else "",
            "category": rider.category if rider else "",
            "team": rider.team if rider else "",
            "sort_time": display_race_time if display_race_time is not None else -1.0,
        }

    def _append_tt_completion_csv(self, row: dict[str, Any]) -> None:
        if not self.tt_export_path:
            return
        self.tt_export_path.parent.mkdir(parents=True, exist_ok=True)
        write_header = not self.tt_export_path.exists() or self.tt_export_path.stat().st_size == 0
        with self.tt_export_path.open("a", newline="") as f:
            writer = csv.writer(f)
            if write_header:
                writer.writerow(["race_time", "bib", "start_time", "stop_time", "elapsed", "last_name", "first_name", "category", "team"])
            writer.writerow([
                row.get("race_time_csv", row.get("race_time", "")),
                row.get("bib", ""),
                row.get("start_time_csv", row.get("start_time", "")),
                row.get("stop_time_csv", row.get("stop_time", "")),
                row.get("elapsed_csv", row.get("elapsed", "")),
                row.get("last_name", ""),
                row.get("first_name", ""),
                row.get("category", ""),
                row.get("team", ""),
            ])

    @staticmethod
    def _clock_to_seconds(value: Any) -> Optional[float]:
        try:
            parsed = parse_clock_to_days(str(value or ""))
        except (TypeError, ValueError):
            return None
        return None if parsed is None else parsed * 86400.0

    @staticmethod
    def _export_read_key(passing: PassingRecord) -> str:
        return (passing.tag or str(passing.bib) or "").strip().upper()

    @staticmethod
    def _export_read_time(passing: PassingRecord) -> Optional[float]:
        if passing.race_time is not None:
            return float(passing.race_time)
        return passing.seen_at.timestamp() if passing.seen_at else None

    def _claim_export_read(
        self,
        seen_by_key: dict[str, list[float]],
        key: str,
        read_time: Optional[float],
        window: float = EXPORT_READ_DEDUP_SECONDS,
    ) -> bool:
        key = key.strip().upper()
        if not key or read_time is None:
            return True
        read_time = float(read_time)
        previous_times = seen_by_key.setdefault(key, [])
        if any(abs(read_time - previous) < window for previous in previous_times):
            return False
        previous_times.append(read_time)
        if len(previous_times) > 20:
            del previous_times[:-20]
        return True

    def _dedupe_export_rows(
        self,
        rows: list[dict[str, Any]],
        key_field: str,
        time_field: str,
        window: float = EXPORT_READ_DEDUP_SECONDS,
    ) -> list[dict[str, Any]]:
        seen_by_key: dict[str, list[float]] = {}
        deduped: list[dict[str, Any]] = []
        for row in rows:
            key = str(row.get(key_field, "")).strip().upper()
            row_time = self._clock_to_seconds(row.get(time_field, ""))
            if self._claim_export_read(seen_by_key, key, row_time, window):
                deduped.append(row)
        return deduped

    def _tt_rider_meta_by_bib(self) -> dict[str, dict[str, str]]:
        category_by_bib: dict[str, str] = {}
        for category in self.config.categories.values():
            display = f"{category.name} (Women)" if category.gender.strip().lower() == "women" else category.name
            for bib in category.numbers.split(","):
                bib_s = bib.strip()
                if bib_s:
                    category_by_bib[bib_s] = display
        return {
            bib: {
                "category": category_by_bib.get(bib, rider.category or ""),
                "team": rider.team or "",
            }
            for bib, rider in self.config.riders_by_bib.items()
        }

    def export_tt_xlsx(self) -> Optional[Path]:
        with self.lock:
            self.tt_tables()
            if not self.is_time_trial() or not self.tt_export_path or not self.tt_export_path.exists():
                return None
            xlsx_path = self.tt_xlsx_path or self.tt_export_path.with_suffix(".xlsx")
            with self.tt_export_path.open(newline="") as f:
                rows = self._dedupe_export_rows(list(csv.DictReader(f)), "bib", "race_time")
            if not rows:
                return None

            rider_meta = self._tt_rider_meta_by_bib()
            wb = Workbook()
            ws = wb.active
            ws.title = "Results"
            wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
            headers = [
                "Pos",
                "Race",
                "BIB",
                "Start",
                "Stop",
                "Penalty",
                "Elapsed",
                "Last",
                "First",
                "Category",
                "Team",
                "Note",
            ]
            ws.append(headers)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
            for cell in ws[1]:
                cell.font = Font(bold=True)

            for row_num, src in enumerate(rows, start=2):
                bib_text = str(src.get("bib", "")).strip()
                meta = rider_meta.get(bib_text, {})
                category = src.get("category", "") or meta.get("category", "")
                team = src.get("team", "") or meta.get("team", "")
                ws[f"A{row_num}"] = (
                    f'=IF(OR(J{row_num}="",G{row_num}=""),"",'
                    f'1+COUNTIFS($J$2:$J$1048576,J{row_num},$G$2:$G$1048576,"<"&G{row_num})'
                    f'+COUNTIFS($J$2:$J$1048576,J{row_num},$G$2:$G$1048576,G{row_num},$C$2:$C$1048576,"<"&C{row_num}))'
                )
                set_time_cell(ws, f"B{row_num}", src.get("race_time", ""), decimals=2)
                ws[f"C{row_num}"] = int(bib_text) if bib_text.isdigit() else bib_text
                set_time_cell(ws, f"D{row_num}", src.get("start_time", ""), decimals=2)
                set_time_cell(ws, f"E{row_num}", src.get("stop_time", ""), decimals=2)
                ws[f"F{row_num}"] = ""
                ws[f"G{row_num}"] = f'=IF(OR(E{row_num}="",D{row_num}=""),"",E{row_num}-D{row_num}+IF(F{row_num}="",0,F{row_num}/86400))'
                ws[f"G{row_num}"].number_format = "[h]:mm:ss.00"
                ws[f"H{row_num}"] = src.get("last_name", "")
                ws[f"I{row_num}"] = src.get("first_name", "")
                ws[f"J{row_num}"] = category
                ws[f"K{row_num}"] = team
                ws[f"L{row_num}"] = ""

            category_values = [str(ws[f"J{row_num}"].value or "") for row_num in range(2, ws.max_row + 1)]
            category_width = max([len("Category"), *(len(v) for v in category_values)], default=len("Category")) + 2
            widths = {
                "A": 5,
                "B": 12,
                "C": 5,
                "D": 12,
                "E": 12,
                "F": 6,
                "G": 12,
                "H": 18,
                "I": 18,
                "J": category_width,
                "K": 22,
                "L": 18,
            }
            for col, width in widths.items():
                ws.column_dimensions[col].width = width
            for row_num in range(2, ws.max_row + 1):
                for col in ("B", "D", "E", "F", "G"):
                    ws[f"{col}{row_num}"].alignment = Alignment(horizontal="center")
                ws[f"C{row_num}"].alignment = Alignment(horizontal="right")

            xlsx_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(xlsx_path)
            return xlsx_path

    def tt_tables(self) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        with self.lock:
            current_race_time = self.current_race_time()
            recorded = self._announcer_recorded_map(current_race_time)
            early_by_bib: dict[str, PassingRecord] = {}
            for passing in self.passings:
                bib = str(passing.bib)
                previous = early_by_bib.get(bib)
                if previous is None:
                    early_by_bib[bib] = passing
                elif passing.race_time is not None and (
                    previous.race_time is None or float(passing.race_time) < float(previous.race_time)
                ):
                    early_by_bib[bib] = passing
            active_rows: list[dict[str, Any]] = []
            completed_rows: list[dict[str, Any]] = []
            is_tt = self.is_time_trial()
            for bib, passing in early_by_bib.items():
                rec = recorded.get(bib)
                rec_t = rec.get("t") if rec else None
                is_completed = False
                if rec_t is not None:
                    if is_tt:
                        is_completed = True
                    elif passing.race_time is not None and float(rec_t) >= float(passing.race_time):
                        is_completed = True
                if is_completed:
                    row = self._tt_row_from_passing(passing, float(rec_t))
                    completed_rows.append(row)
                    export_key = f"{bib}:{row['early_time']}"
                    if export_key not in self.tt_exported_keys:
                        export_read_key = self._export_read_key(passing)
                        export_read_time = self._export_read_time(passing)
                        if self._claim_export_read(self.tt_exported_read_times, export_read_key, export_read_time):
                            self._append_tt_completion_csv(row)
                        self.tt_exported_keys.add(export_key)
                else:
                    active_rows.append(self._tt_row_from_passing(passing, None))
                    if bib not in self.tt_diag_logged:
                        result = self.results_by_bib.get(str(bib))
                        if result:
                            race_times = self._race_times(result)
                            interp = self._interp_flags(result)
                            logging.getLogger("lapsrv.tt").warning(
                                "tt unresolved bib=%s status=%s startTime=%r raceTimes_tail=%s interp_tail=%s rec=%s early=%s",
                                bib,
                                result.get("status"),
                                result.get("startTime"),
                                race_times[-3:],
                                interp[-3:],
                                rec,
                                passing.race_time,
                            )
                        else:
                            logging.getLogger("lapsrv.tt").warning(
                                "tt unresolved bib=%s no announcer result early=%s",
                                bib,
                                passing.race_time,
                            )
                        self.tt_diag_logged.add(bib)
        active_rows.sort(key=lambda r: r.get("sort_time") or -1.0, reverse=True)
        completed_rows.sort(key=lambda r: r.get("sort_time") or -1.0, reverse=True)
        for rows in (active_rows, completed_rows):
            for row in rows:
                row.pop("sort_time", None)
        return active_rows, completed_rows

    def _road_export_row_from_match(self, passing: PassingRecord, recorded: dict[str, Any]) -> dict[str, Any]:
        rider = self.config.riders_by_bib.get(str(passing.bib))
        result = self.results_by_bib.get(str(passing.bib), {})
        category = self._effective_category_name(rider, result) if rider else passing.category
        early_time = passing.race_time
        crossmgr_time = float(recorded["t"]) if recorded.get("t") is not None else None
        time_delta = None if early_time is None or crossmgr_time is None else max(0.0, crossmgr_time - early_time)
        early_pos = passing.category_rank
        crossmgr_pos = self._rank_for_rider(rider) if rider else None
        pos_change = None
        if early_pos is not None and crossmgr_pos is not None:
            pos_change = int(early_pos) - int(crossmgr_pos)
        return {
            "early_time": format_clock_ms(early_time),
            "crossmgr_time": format_clock_ms(crossmgr_time),
            "time_delta": format_clock_ms(time_delta),
            "bib": str(passing.bib),
            "last_name": rider.last_name if rider else "",
            "first_name": rider.first_name if rider else "",
            "category": category,
            "team": rider.team if rider else passing.team,
            "lap": str(recorded.get("lap") or passing.lap or ""),
            "early_pos": str(early_pos or ""),
            "crossmgr_pos": str(crossmgr_pos or ""),
            "pos_change": "" if pos_change is None else str(pos_change),
        }

    def _append_road_passing_csv(self, row: dict[str, Any]) -> None:
        if not self.road_export_path:
            return
        self.road_export_path.parent.mkdir(parents=True, exist_ok=True)
        write_header = not self.road_export_path.exists() or self.road_export_path.stat().st_size == 0
        fields = [
            "early_time",
            "crossmgr_time",
            "time_delta",
            "bib",
            "last_name",
            "first_name",
            "category",
            "team",
            "lap",
            "early_pos",
            "crossmgr_pos",
            "pos_change",
        ]
        with self.road_export_path.open("a", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            if write_header:
                writer.writeheader()
            writer.writerow({field: row.get(field, "") for field in fields})

    def update_road_exports(self) -> None:
        with self.lock:
            if self.is_time_trial() or not self.road_export_path:
                return
            recorded_by_bib = self._road_recorded_events_by_bib()
            used_recorded_keys: set[str] = set()
            for passing in sorted(self.road_pending_passings, key=lambda p: p.race_time if p.race_time is not None else -1.0):
                if passing.race_time is None or not passing.bib or passing.bib in {"?", ""}:
                    continue
                early_key = f"{passing.bib}:{format_clock_ms(passing.race_time)}"
                if early_key in self.road_exported_keys:
                    continue
                recorded_events = recorded_by_bib.get(str(passing.bib), [])
                match: Optional[dict[str, Any]] = None
                for recorded in recorded_events:
                    rec_t = recorded.get("t")
                    if rec_t is None or float(rec_t) + 0.001 < float(passing.race_time):
                        continue
                    recorded_key = f"{passing.bib}:{recorded.get('lap')}:{format_clock_ms(float(rec_t))}"
                    if recorded_key in used_recorded_keys or recorded_key in self.road_exported_recorded_keys:
                        continue
                    match = recorded
                    used_recorded_keys.add(recorded_key)
                    break
                if not match:
                    continue
                export_read_key = self._export_read_key(passing)
                export_read_time = self._export_read_time(passing)
                if not self._claim_export_read(self.road_exported_read_times, export_read_key, export_read_time):
                    self.road_exported_keys.add(early_key)
                    continue
                self._append_road_passing_csv(self._road_export_row_from_match(passing, match))
                self.road_exported_keys.add(early_key)
                self.road_exported_recorded_keys.add(f"{passing.bib}:{match.get('lap')}:{format_clock_ms(float(match.get('t')))}")

    def export_road_xlsx(self) -> Optional[Path]:
        with self.lock:
            self.update_road_exports()
            if self.is_time_trial() or not self.road_export_path or not self.road_export_path.exists():
                return None
            xlsx_path = self.road_xlsx_path or self.road_export_path.with_suffix(".xlsx")
            with self.road_export_path.open(newline="") as f:
                rows = self._dedupe_export_rows(list(csv.DictReader(f)), "bib", "early_time")
            if not rows:
                return None

            wb = Workbook()
            ws = wb.active
            ws.title = "Road Passings"
            wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
            headers = [
                "Early",
                "CrossMgr",
                "Delta",
                "BIB",
                "Last",
                "First",
                "Category",
                "Team",
                "Lap",
                "Early Pos",
                "CrossMgr Pos",
                "Pos +/-",
                "Note",
            ]
            ws.append(headers)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
            for cell in ws[1]:
                cell.font = Font(bold=True)

            for row_num, src in enumerate(rows, start=2):
                set_time_cell(ws, f"A{row_num}", src.get("early_time", ""), decimals=3)
                set_time_cell(ws, f"B{row_num}", src.get("crossmgr_time", ""), decimals=3)
                set_time_cell(ws, f"C{row_num}", src.get("time_delta", ""), decimals=3)
                bib = str(src.get("bib", "")).strip()
                ws[f"D{row_num}"] = int(bib) if bib.isdigit() else bib
                ws[f"E{row_num}"] = src.get("last_name", "")
                ws[f"F{row_num}"] = src.get("first_name", "")
                ws[f"G{row_num}"] = src.get("category", "")
                ws[f"H{row_num}"] = src.get("team", "")
                lap = str(src.get("lap", "")).strip()
                ws[f"I{row_num}"] = int(lap) if lap.isdigit() else lap
                for col, field in (("J", "early_pos"), ("K", "crossmgr_pos"), ("L", "pos_change")):
                    value = str(src.get(field, "")).strip()
                    ws[f"{col}{row_num}"] = int(value) if value.lstrip("-").isdigit() else value
                ws[f"M{row_num}"] = ""

            widths = {
                "A": 13,
                "B": 13,
                "C": 13,
                "D": 5,
                "E": 18,
                "F": 18,
                "G": max(12, min(32, max(len(str(r.get("category", ""))) for r in rows) + 2)),
                "H": 24,
                "I": 5,
                "J": 9,
                "K": 11,
                "L": 8,
                "M": 22,
            }
            for col, width in widths.items():
                ws.column_dimensions[col].width = width
            for row_num in range(2, ws.max_row + 1):
                for col in ("A", "B", "C", "D", "I", "J", "K", "L"):
                    ws[f"{col}{row_num}"].alignment = Alignment(horizontal="center")

            xlsx_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(xlsx_path)
            return xlsx_path

    def recent_group_table(self) -> dict[str, Any]:
        recent, _past = self._build_group_tables()
        return recent

    def past_group_table(self) -> dict[str, Any]:
        _recent, past = self._build_group_tables()
        return past

    def category_passings(self, category: str) -> list[dict[str, Any]]:
        with self.lock:
            return [p.as_dict() for p in self.passings if p.category == category]

    def configure_ip_health(self, addresses: list[str]) -> None:
        with self.lock:
            self.ip_health = {addr: self.ip_health.get(addr, IPHealthStatus(address=addr)) for addr in addresses}

    def update_ip_health(self, address: str, ok: bool, rtt_ms: Optional[float] = None) -> None:
        with self.lock:
            status = self.ip_health.setdefault(address, IPHealthStatus(address=address))
            status.last_probe_at = dt.datetime.now().astimezone()
            status.last_probe_ok = ok
            if ok:
                status.last_ok_at = status.last_probe_at
                status.last_rtt_ms = rtt_ms

    def ip_health_snapshot(self) -> list[dict[str, Any]]:
        with self.lock:
            return [self.ip_health[address].snapshot() for address in self.ip_health]

    def has_ip_health_issue(self) -> bool:
        return any(item.get("state") in {"missed", "down"} for item in self.ip_health_snapshot())

    def set_warning_enabled(self, enabled: bool) -> None:
        with self.lock:
            self.warning_enabled = bool(enabled)

    def snapshot(self) -> dict[str, Any]:
        with self.lock:
            current_race_time = self.current_race_time()
            is_time_trial = self.is_time_trial()
            if is_time_trial:
                tt_active_rows, tt_completed_rows = self.tt_tables()
            else:
                self.update_road_exports()
                tt_active_rows, tt_completed_rows = [], []
            return {
                "event_name": self.config.event_name,
                "event_date": self.config.event_date,
                "timezone": self.config.timezone,
                "current_race_time": format_clock(current_race_time),
                "cur_race_time": format_clock(current_race_time),
                "is_time_trial": is_time_trial,
                "tt_active_rows": tt_active_rows,
                "tt_completed_rows": tt_completed_rows,
                "tt_csv_path": str(self.tt_export_csv_path() or ""),
                "road_csv_path": str(self.road_export_csv_path() or ""),
                "recent_group_table": self.recent_group_table(),
                "past_group_table": self.past_group_table(),
                "passings": [p.as_dict() for p in self.passings],
                "readers": dict(self.reader_status),
                "reader_clock_diff": {k: round(v, 3) for k, v in self.reader_clock_diff.items()},
                "last_announcer_update": as_utc_iso(self.last_announcer_update) if self.last_announcer_update else "",
                "last_lapcounter_update": as_utc_iso(self.last_lapcounter_update) if self.last_lapcounter_update else "",
                "latest_lap_refresh": self.latest_lap_refresh,
                "subsystems": {k: self._effective_subsystem_state(k).snapshot() for k in self.subsystems},
                "ip_health": self.ip_health_snapshot(),
                "warning_enabled": self.warning_enabled,
            }


class JChipServer:
    def __init__(self, state: SharedState, host: str, port: int) -> None:
        self.state = state
        self.host = host
        self.port = port
        self.server: Optional[asyncio.AbstractServer] = None
        self.logger = logging.getLogger("lapsrv.jchip")
        self.date_today = dt.date.today()
        self.last_timestamp: Optional[dt.datetime] = None
        self.same_count = 0
        self.client_tasks: set[asyncio.Task[Any]] = set()
        self.client_writers: set[asyncio.StreamWriter] = set()

    async def start(self) -> None:
        self.server = await asyncio.start_server(self.handle_client, self.host, self.port)
        socknames = ", ".join(str(sock.getsockname()) for sock in (self.server.sockets or []))
        self.logger.info("JChip listener started on %s", socknames)
        self.state.set_subsystem_status("rfid", "waiting", endpoint=socknames)

    async def stop(self) -> None:
        if self.server:
            self.server.close()
            await self.server.wait_closed()
            self.server = None
        for writer in list(self.client_writers):
            writer.close()
        for writer in list(self.client_writers):
            try:
                await asyncio.wait_for(writer.wait_closed(), timeout=1.0)
            except Exception:
                pass
        if self.client_tasks:
            for task in list(self.client_tasks):
                task.cancel()
            try:
                await asyncio.wait_for(asyncio.gather(*list(self.client_tasks), return_exceptions=True), timeout=1.0)
            except Exception:
                pass

    def _parse_time(self, t_str: str, day: int = 0) -> dt.datetime:
        hh, mm, ss = t_str.split(":")
        base = dt.datetime.combine(self.date_today, dt.time()) + dt.timedelta(
            seconds=(float(hh) * 3600.0) + (float(mm) * 60.0) + float(ss)
        )
        timestamp = base + dt.timedelta(days=day)
        if self.last_timestamp is None:
            self.last_timestamp = timestamp - dt.timedelta(microseconds=10)
        if timestamp == self.last_timestamp:
            self.same_count += 1
        else:
            self.same_count = 0
            self.last_timestamp = timestamp
        if self.same_count:
            timestamp += dt.timedelta(microseconds=10 * self.same_count)
        return timestamp.astimezone()

    def _parse_tag_line(self, line: str) -> tuple[str, dt.datetime]:
        i_space = line.find(" ")
        if i_space < 0:
            raise ValueError(f"invalid D record: {line!r}")
        tag = line[2:i_space].strip().upper()
        i_colon = line.find(":")
        if i_colon < 2:
            raise ValueError(f"missing time in D record: {line!r}")
        match = RE_TIME.match(line[i_colon - 2 :])
        if not match:
            raise ValueError(f"invalid time in D record: {line!r}")
        t_str = match.group(0)
        day = 0
        i_second_field = line.find(" ", i_colon)
        if i_second_field >= 0 and i_second_field + 2 < len(line):
            try:
                day = int(line[i_second_field + 2 : i_second_field + 3])
            except ValueError:
                day = 0
        if "date=" in line:
            i_date = line.index("date=") + 5
            year = int(line[i_date : i_date + 4])
            month = int(line[i_date + 4 : i_date + 6])
            day_of_month = int(line[i_date + 6 : i_date + 8])
            tag_date = dt.date(year, month, day_of_month)
            day = (tag_date - self.date_today).days
        return tag, self._parse_time(t_str, day)

    async def handle_client(self, reader: asyncio.StreamReader, writer: asyncio.StreamWriter) -> None:
        task = asyncio.current_task()
        if task is not None:
            self.client_tasks.add(task)
        self.client_writers.add(writer)
        addr = writer.get_extra_info("peername")
        reader_name = f"{addr}"
        time_adjust = 0.0
        self.logger.info("reader connected: %s", addr)
        self.state.reset_for_reader_reconnect(reader_name)
        try:
            while not reader.at_eof():
                raw = await reader.readuntil(CR_BYTES)
                line = raw[:-1].decode(errors="ignore").strip()
                if not line:
                    continue
                self.logger.warning("reader rx: %s %s", reader_name, line)
                if line.startswith("N"):
                    reader_name = line[5:].strip() or reader_name
                    self.logger.info("reader identified: %s", reader_name)
                    self.state.set_reader_status(reader_name, "connected")
                    writer.write(f"GT{CR}".encode())
                    await writer.drain()
                    self.logger.info("sent GT to reader: %s", reader_name)
                elif line.startswith("GT"):
                    now = dt.datetime.now()
                    hh = int(line[3:5])
                    mm = int(line[5:7])
                    ss = int(line[7:9])
                    hs = int(line[9:11])
                    if "date=" in line:
                        i_date = line.index("date=") + 5
                        jchip_now = dt.datetime(
                            int(line[i_date : i_date + 4]),
                            int(line[i_date + 4 : i_date + 6]),
                            int(line[i_date + 6 : i_date + 8]),
                            hh,
                            mm,
                            ss,
                            hs * 10000,
                        )
                        self.date_today = jchip_now.date()
                    else:
                        jchip_now = dt.datetime.combine(now.date(), dt.time(hh, mm, ss, hs * 10000))
                    time_adjust = (now - jchip_now).total_seconds()
                    self.state.set_reader_clock_diff(reader_name, time_adjust)
                    writer.write(f"S0000{CR}".encode())
                    await writer.drain()
                    self.logger.info("sent S0000 to reader: %s", reader_name)
                elif line.startswith("D"):
                    tag, seen_at = self._parse_tag_line(line)
                    seen_at = seen_at + dt.timedelta(seconds=time_adjust)
                    record = self.state.record_passing(tag, seen_at.astimezone(), reader_name)
                    if record is not None:
                        self.logger.info(
                            "passing: reader=%s tag=%s bib=%s cat=%s time=%s",
                            reader_name,
                            record.tag,
                            record.bib,
                            record.category,
                            format_clock(record.race_time),
                        )
                else:
                    self.logger.debug("ignored reader line from %s: %s", reader_name, line)
        except (asyncio.IncompleteReadError, ConnectionResetError):
            pass
        except asyncio.CancelledError:
            pass
        except Exception:
            self.logger.exception("reader handler failed for %s", reader_name)
            self.state.set_subsystem_status("rfid", "error", detail=reader_name)
        finally:
            self.state.set_reader_status(reader_name, "disconnected")
            self.client_writers.discard(writer)
            writer.close()
            try:
                await asyncio.wait_for(writer.wait_closed(), timeout=1.0)
            except Exception:
                pass
            if task is not None:
                self.client_tasks.discard(task)
            self.logger.info("reader disconnected: %s", reader_name)


class CrossMgrClient:
    def __init__(self, state: SharedState, host: str, port: int) -> None:
        self.state = state
        self.host = host
        self.port = port
        self.logger = logging.getLogger("lapsrv.crossmgr")
        self.stop_event = asyncio.Event()
        self._announcer_ws: Any = None
        self._lapcounter_ws: Any = None

    async def run(self) -> None:
        await asyncio.gather(
            self._run_announcer(),
            self._run_lapcounter(),
        )

    async def stop(self) -> None:
        self.stop_event.set()
        for ws in (self._announcer_ws, self._lapcounter_ws):
            if ws is not None:
                try:
                    await asyncio.wait_for(ws.close(), timeout=1.0)
                except Exception:
                    pass

    async def _run_announcer(self) -> None:
        url = f"ws://{self.host}:{self.port + 1}/"
        while not self.stop_event.is_set():
            try:
                self.logger.info("connecting announcer websocket %s", url)
                self.state.set_subsystem_status("announcer", "waiting", endpoint=self.host)
                async with websockets.connect(url, open_timeout=5, ping_interval=None) as ws:
                    self._announcer_ws = ws
                    self.state.set_subsystem_status("announcer", "connected", endpoint=self.host)
                    await ws.send(json.dumps({"cmd": "send_baseline", "raceName": "CurrentResults"}))
                    async for raw in ws:
                        if self.stop_event.is_set():
                            break
                        msg = json.loads(raw)
                        if msg.get("cmd") in {"baseline", "ram"}:
                            self.state.update_announcer(msg)
                self._announcer_ws = None
            except asyncio.CancelledError:
                self._announcer_ws = None
                raise
            except Exception as exc:
                self.logger.warning("announcer websocket error: %s", exc)
                self.state.set_subsystem_status("announcer", "error", endpoint=self.host, detail=str(exc))
                await asyncio.sleep(5)

    async def _run_lapcounter(self) -> None:
        url = f"ws://{self.host}:{self.port + 2}/"
        while not self.stop_event.is_set():
            try:
                self.logger.info("connecting lapcounter websocket %s", url)
                self.state.set_subsystem_status("lapcounter", "waiting", endpoint=self.host)
                async with websockets.connect(url, open_timeout=5, ping_interval=None) as ws:
                    self._lapcounter_ws = ws
                    self.state.set_subsystem_status("lapcounter", "connected", endpoint=self.host)
                    async for raw in ws:
                        if self.stop_event.is_set():
                            break
                        msg = json.loads(raw)
                        if msg.get("cmd") == "refresh":
                            self.state.update_lapcounter(msg)
                self._lapcounter_ws = None
            except asyncio.CancelledError:
                self._lapcounter_ws = None
                raise
            except Exception as exc:
                self.logger.warning("lapcounter websocket error: %s", exc)
                self.state.set_subsystem_status("lapcounter", "error", endpoint=self.host, detail=str(exc))
                await asyncio.sleep(5)


class PingMonitor:
    def __init__(self, state: SharedState, addresses: list[str]) -> None:
        self.state = state
        self.addresses = addresses[:3]
        self.logger = logging.getLogger("lapsrv.ping")
        self.stop_event = asyncio.Event()
        self.ping_path = which("ping") or "ping"
        self.player_cmd = pick_audio_player()
        self.player_warned = False
        self.last_alert_at = 0.0
        self.had_issue = False
        pcm_a, sample_rate_a = build_tone_pcm(frequency=988.0, duration=0.12)
        pcm_b, sample_rate_b = build_tone_pcm(frequency=1318.0, duration=0.12)
        self.alert_wav_a = build_wave_bytes(pcm_a, sample_rate_a)
        self.alert_wav_b = build_wave_bytes(pcm_b, sample_rate_b)
        self.state.configure_ip_health(self.addresses)

    async def stop(self) -> None:
        self.stop_event.set()

    async def _play_alert_tone(self) -> None:
        if self.player_cmd is None:
            if not self.player_warned:
                self.logger.warning("no audio player found for ping alert; tried paplay, aplay, ffplay, play")
                self.player_warned = True
            return
        try:
            for wav_bytes in (self.alert_wav_a, self.alert_wav_b):
                proc = await asyncio.create_subprocess_exec(
                    *self.player_cmd,
                    stdin=asyncio.subprocess.PIPE,
                    stdout=asyncio.subprocess.DEVNULL,
                    stderr=asyncio.subprocess.DEVNULL,
                )
                await proc.communicate(wav_bytes)
                await asyncio.sleep(0.08)
        except Exception as exc:
            self.logger.warning("ping alert tone failed: %s", exc)

    async def _ping_once(self, address: str) -> None:
        if not address:
            return
        try:
            proc = await asyncio.create_subprocess_exec(
                self.ping_path, "-c", "1", "-W", "1", address,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
            )
            try:
                stdout, _stderr = await asyncio.wait_for(proc.communicate(), timeout=2.0)
            except asyncio.TimeoutError:
                proc.kill()
                await proc.communicate()
                self.state.update_ip_health(address, False)
                return
            ok = proc.returncode == 0
            rtt_ms = None
            if ok:
                match = re.search(rb"time=([0-9]+(?:\.[0-9]+)?)", stdout or b"")
                if match:
                    try:
                        rtt_ms = float(match.group(1))
                    except ValueError:
                        rtt_ms = None
            self.state.update_ip_health(address, ok, rtt_ms)
        except Exception as exc:
            self.logger.debug("ping failed for %s: %s", address, exc)
            self.state.update_ip_health(address, False)

    async def run(self) -> None:
        if not self.addresses:
            return
        while not self.stop_event.is_set():
            await asyncio.gather(*(self._ping_once(address) for address in self.addresses), return_exceptions=True)
            has_issue = self.state.has_ip_health_issue()
            if has_issue and not self.had_issue:
                self.state.set_warning_enabled(True)
                self.last_alert_at = 0.0
            self.had_issue = has_issue
            if self.state.warning_enabled and has_issue and (time.time() - self.last_alert_at) >= 5.0:
                await self._play_alert_tone()
                self.last_alert_at = time.time()
            try:
                await asyncio.wait_for(self.stop_event.wait(), timeout=1.0)
            except asyncio.TimeoutError:
                pass


class WebServer:
    def __init__(self, state: SharedState, host: str, port: int) -> None:
        self.state = state
        self.host = host
        self.port = port
        self.app = web.Application()
        self.runner: Optional[web.AppRunner] = None
        self.site: Optional[web.TCPSite] = None
        self.logger = logging.getLogger("lapsrv.web")
        self._routes()

    def _routes(self) -> None:
        self.app.router.add_get("/", self.handle_index)
        self.app.router.add_get("/category/{name}", self.handle_category)
        self.app.router.add_get("/api/state", self.handle_state)
        self.app.router.add_get("/api/category/{name}", self.handle_category_state)
        self.app.router.add_post("/api/warning", self.handle_warning)

    async def start(self) -> None:
        self.runner = web.AppRunner(self.app)
        await self.runner.setup()
        self.site = web.TCPSite(self.runner, self.host, self.port)
        await self.site.start()
        self.logger.info("web server started on http://%s:%s", self.host, self.port)

    async def stop(self) -> None:
        if self.runner:
            await self.runner.cleanup()

    async def handle_state(self, request: web.Request) -> web.Response:
        return web.json_response(self.state.snapshot())

    async def handle_category_state(self, request: web.Request) -> web.Response:
        category = request.match_info["name"]
        snapshot = self.state.snapshot()
        snapshot["passings"] = self.state.category_passings(category)
        snapshot["selected_category"] = category
        return web.json_response(snapshot)

    async def handle_warning(self, request: web.Request) -> web.Response:
        payload = await request.json()
        self.state.set_warning_enabled(bool(payload.get("enabled", True)))
        return web.json_response({"ok": True, "warning_enabled": self.state.snapshot().get("warning_enabled", True)})

    async def handle_index(self, request: web.Request) -> web.Response:
        return web.Response(text=self._html(category=None), content_type="text/html")

    async def handle_category(self, request: web.Request) -> web.Response:
        return web.Response(text=self._html(category=request.match_info["name"]), content_type="text/html")

    def _html(self, category: Optional[str]) -> str:
        category_js = json.dumps(category)
        api_path = "/api/state" if category is None else f"/api/category/{category}"
        return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>lapsrv</title>
<style>
:root {{
  --bg: #f8fafc;
  --panel: #ffffff;
  --panel-2: #e5eef7;
  --fg: #0f172a;
  --muted: #475569;
  --accent: #b45309;
  --ok: #15803d;
  --border: #cbd5e1;
}}
* {{ box-sizing: border-box; }}
body {{
  margin: 0;
  font-family: "Trebuchet MS", "Segoe UI", sans-serif;
  background: linear-gradient(180deg, #fefefe 0%, #eef6ff 100%);
  color: var(--fg);
}}
body.compact-mode .default-view, body.compact-mode .tt-view {{ display: none; }}
body.tt-mode .default-view, body.tt-mode .compact-view {{ display: none; }}
body:not(.compact-mode):not(.tt-mode) .compact-view, body:not(.compact-mode):not(.tt-mode) .tt-view {{ display: none; }}
body.compact-mode {{
  overflow: hidden;
}}
body.tt-mode {{
  overflow: hidden;
}}
main {{
  max-width: 1400px;
  margin: 0 auto;
  padding: 16px;
  min-height: 100vh;
  display: flex;
  flex-direction: column;
}}
body.compact-mode main {{
  height: 100vh;
  min-height: 100vh;
  overflow: hidden;
}}
.title {{
  display: grid;
  grid-template-columns: auto 1fr auto;
  align-items: flex-start;
  gap: 8px;
  margin-bottom: 2px;
}}
.title-left {{
  justify-self: start;
}}
.title-center {{
  justify-self: center;
  text-align: center;
}}
.compact-race-time {{
  font-size: 1.2rem;
  font-weight: 700;
  line-height: 1.1;
  color: var(--fg);
  white-space: nowrap;
}}
.meta {{ color: var(--muted); font-size: 0.95rem; }}
.statusbar {{
  display: flex;
  gap: 10px;
  flex-wrap: wrap;
  align-items: center;
}}
.toolbar {{
  display: flex;
  gap: 8px;
  align-items: flex-start;
  flex-wrap: nowrap;
  margin-top: 0;
}}
.tt-indicator {{
  border: 1px solid rgba(21,128,61,.35);
  background: rgba(21,128,61,.12);
  color: #166534;
  border-radius: 999px;
  padding: 5px 10px;
  font-size: 0.78rem;
  font-weight: 700;
  width: 52px;
  text-align: center;
  white-space: nowrap;
  line-height: 1.05;
}}
.view-toggle, .audio-toggle, .warning-toggle {{
  border: 1px solid rgba(250,204,21,.45);
  background: rgba(250,204,21,.12);
  color: #92400e;
  border-radius: 999px;
  padding: 5px 10px;
  font-size: 0.78rem;
  font-weight: 700;
  cursor: pointer;
  width: 94px;
  text-align: center;
  white-space: nowrap;
  line-height: 1.05;
}}
.compact-toggle {{
  border: 1px solid rgba(148,163,184,.45);
  background: rgba(148,163,184,.12);
  color: #334155;
  border-radius: 999px;
  padding: 5px 10px;
  font-size: 0.78rem;
  font-weight: 700;
  cursor: pointer;
  width: 98px;
  text-align: center;
  white-space: nowrap;
  line-height: 1.05;
}}
.compact-toggle.active {{
  background: rgba(21,128,61,.12);
  border-color: rgba(21,128,61,.35);
  color: #166534;
}}
.audio-toggle.active, .warning-toggle.active {{
  background: rgba(21,128,61,.12);
  border-color: rgba(21,128,61,.35);
  color: #166534;
}}
.ip-healthbar {{
  display: flex;
  gap: 4px;
  flex-wrap: wrap;
  margin: -2px 0 8px;
  padding-left: 2px;
}}
.ip-pill {{
  border-radius: 2px;
  padding: 3px 6px;
  font-weight: 700;
  font-size: 0.78rem;
  line-height: 1.0;
  border: 1px solid #94a3b8;
}}
.ip-pill.ok {{
  background: rgba(34,197,94,.16);
  color: #166534;
}}
.ip-pill.missed {{
  background: rgba(245,158,11,.16);
  color: #92400e;
}}
.ip-pill.down {{
  background: #fbcfe8;
  color: #9d174d;
}}
.badge {{
  display: inline-flex;
  gap: 6px;
  align-items: center;
  border-radius: 999px;
  padding: 6px 10px;
  font-size: 0.92rem;
  font-weight: 700;
  border: 1px solid transparent;
}}
.state-connected {{
  background: rgba(34,197,94,.12);
  color: #166534;
  border-color: rgba(34,197,94,.24);
}}
.state-waiting, .state-timedout {{
  background: rgba(245,158,11,.14);
  color: #92400e;
  border-color: rgba(245,158,11,.28);
}}
.state-error {{
  background: rgba(239,68,68,.12);
  color: #991b1b;
  border-color: rgba(239,68,68,.24);
}}
.group-wrap {{
  margin: 16px 0 18px;
  background: rgba(255,255,255,.96);
  border: 1px solid var(--border);
  border-radius: 14px;
  overflow: hidden;
  display: flex;
  flex-direction: column;
}}
.group-title {{
  padding: 10px 14px;
  font-size: 1.05rem;
  font-weight: 700;
  color: var(--muted);
  background: rgba(229,238,247,.98);
  border-bottom: 1px solid var(--border);
}}
.group-rows {{
  overflow-y: auto;
  user-select: text;
  -webkit-user-select: text;
  background: rgba(255,255,255,.98);
}}
.recent-group-rows {{
  height: 48vh;
}}
body.compact-mode .recent-group-rows {{
  height: auto;
  flex: 1 1 auto;
  min-height: 0;
  overflow-y: auto;
}}
.past-group-rows {{
  height: 26vh;
}}
.group-table, .small-table, .tt-table {{
  background: rgba(255,255,255,.98);
}}
.group-table tbody tr, .small-table tbody tr, .tt-table tbody tr {{
  background: rgba(255,255,255,.98);
}}
.tt-scroll, .table-wrap {{
  background: rgba(255,255,255,.98);
}}
.group-table {{
  width: 100%;
  border-collapse: collapse;
  table-layout: fixed;
}}
.group-table th, .group-table td {{
  padding: 8px 10px;
  border-bottom: 1px solid rgba(203,213,225,.95);
  font-size: 0.98rem;
  vertical-align: top;
}}
.group-table th {{
  position: sticky;
  top: 0;
  background: rgba(229,238,247,.98);
  color: #0f172a;
  z-index: 1;
}}
.group-table thead tr:nth-child(2) th {{
  top: 42px;
}}
.group-cell.lead, .group-header-status.bell {{
  background: #facc15;
  color: #111827;
  font-weight: 700;
}}
.group-cell-content {{
  display: flex;
  align-items: baseline;
  justify-content: space-between;
  gap: 8px;
  width: 100%;
}}
.group-cell-note {{
  min-width: 0;
}}
.group-cell-count {{
  margin-left: auto;
  text-align: right;
  min-width: 1ch;
}}
.group-cell.lapped {{
  background: rgba(226, 232, 240, 0.16);
  color: #64748b;
}}
.group-cell.lapped .group-cell-content {{
  justify-content: flex-end;
}}
.group-header-status {{
  color: #e5e7eb;
  font-weight: 700;
}}
.group-empty {{ color: #6b7280; }}
.compact-wrap {{
  margin: 16px 0 0;
  background: rgba(255,255,255,.96);
  border: 1px solid var(--border);
  border-radius: 14px;
  overflow: hidden;
  flex: 1 1 auto;
  min-height: 0;
  display: flex;
  flex-direction: column;
}}
.compact-table {{
  width: 100%;
  border-collapse: collapse;
  table-layout: fixed;
}}
.compact-table th:nth-child(1), .compact-table td:nth-child(1) {{
  width: 72px;
  white-space: nowrap;
  text-align: center;
}}
.compact-table th:nth-child(3), .compact-table td:nth-child(3) {{
  width: 88px;
  white-space: nowrap;
  text-align: center;
}}
.compact-table th:nth-child(4), .compact-table td:nth-child(4) {{
  width: 76px;
  white-space: nowrap;
  text-align: center;
}}
.compact-table-separate th:nth-child(1), .compact-table-separate td:nth-child(1) {{
  width: 56px;
  white-space: nowrap;
  text-align: center;
}}
.compact-table-separate th:nth-child(2), .compact-table-separate td:nth-child(2) {{
  width: auto;
  text-align: left;
}}
.compact-table-separate th:nth-child(3), .compact-table-separate td:nth-child(3) {{
  width: 92px;
  white-space: nowrap;
  text-align: center;
}}
.compact-table-separate th:nth-child(4), .compact-table-separate td:nth-child(4) {{
  width: 76px;
  white-space: nowrap;
  text-align: center;
}}
.compact-table-separate th:nth-child(5), .compact-table-separate td:nth-child(5) {{
  width: 84px;
  white-space: nowrap;
  text-align: right;
}}
.compact-table-separate .compact-group-cell {{
  font-size: 1.95rem;
  line-height: 1.0;
}}
.compact-table th, .compact-table td {{
  padding: 8px 10px;
  border-bottom: 1px solid rgba(203,213,225,.95);
  font-size: 1.02rem;
  vertical-align: top;
  background: rgba(255,255,255,.98);
}}
.compact-table th {{
  background: rgba(229,238,247,.98);
  position: sticky;
  top: 0;
  z-index: 1;
}}
.compact-group-cell {{
  font-weight: 700;
  letter-spacing: 0.01em;
  font-size: 2.3rem;
  line-height: 1.1;
}}
.compact-group-lead td, .compact-group-lead .compact-group-cell {{
  background: #facc15;
  color: #111827;
}}
.compact-group-bell td, .compact-group-bell .compact-group-cell {{
  background: #fbcfe8;
  color: #111827;
}}
.compact-group-lapped td, .compact-group-lapped .compact-group-cell {{
  background: rgba(226, 232, 240, 0.55);
  color: #64748b;
}}
.compact-separator td {{
  padding: 0;
  border-bottom: none;
  height: 14px;
  background: transparent;
}}
table {{
  width: 100%;
  border-collapse: collapse;
  background: rgba(17,24,39,.92);
  border: 1px solid var(--border);
  border-radius: 14px;
  overflow: hidden;
}}
th, td {{
  padding: 10px 8px;
  border-bottom: 1px solid var(--border);
  text-align: left;
  font-size: 0.95rem;
}}
th {{
  background: rgba(229,238,247,.98);
  color: #0f172a;
  position: sticky;
  top: 0;
}}
tr:last-child td {{ border-bottom: none; }}
a {{ color: #92400e; text-decoration: none; }}
.note {{ color: var(--accent); font-weight: 700; }}
.bell {{ color: var(--ok); font-size: 1.2rem; }}
.table-wrap {{
  max-height: 24vh;
  overflow-y: auto;
  border-radius: 14px;
}}
.small-table th, .small-table td {{
  padding: 7px 6px;
  font-size: 0.8rem;
}}
body.is-mobile .recent-group-rows {{
  height: 52vh;
}}
body.is-mobile.compact-mode .recent-group-rows {{
  height: auto;
}}
body.is-mobile .past-group-rows {{
  height: 22vh;
}}
body.is-mobile .compact-table th:nth-child(1), body.is-mobile .compact-table td:nth-child(1) {{
  width: 62px;
}}
body.is-mobile .compact-table th:nth-child(3), body.is-mobile .compact-table td:nth-child(3) {{
  width: 74px;
}}
body.is-mobile .compact-table th:nth-child(4), body.is-mobile .compact-table td:nth-child(4) {{
  width: 64px;
}}
body.is-mobile .compact-table-separate th:nth-child(1), body.is-mobile .compact-table-separate td:nth-child(1) {{
  width: 48px;
}}
body.is-mobile .compact-table-separate th:nth-child(3), body.is-mobile .compact-table-separate td:nth-child(3) {{
  width: 72px;
}}
body.is-mobile .compact-table-separate th:nth-child(4), body.is-mobile .compact-table-separate td:nth-child(4) {{
  width: 60px;
}}
body.is-mobile .compact-table-separate th:nth-child(5), body.is-mobile .compact-table-separate td:nth-child(5) {{
  width: 72px;
}}
body.is-mobile .compact-table-separate .compact-group-cell {{
  font-size: 1.8rem;
}}
body.is-mobile .compact-table th, body.is-mobile .compact-table td {{
  font-size: 0.95rem;
  padding: 8px 6px;
}}
body.is-mobile .compact-group-cell {{
  font-size: 2.1rem;
  line-height: 1.05;
}}
body.is-mobile th:nth-child(6), body.is-mobile td:nth-child(6),
body.is-mobile th:nth-child(8), body.is-mobile td:nth-child(8) {{ display: none; }}
body.is-desktop-portrait {{
  font-size: 200%;
}}
body.is-desktop-portrait main {{
  max-width: none;
  padding: 20px;
}}
body.is-desktop-portrait #event {{
  font-size: 1.1rem !important;
}}
body.is-desktop-portrait .warning-toggle,
body.is-desktop-portrait .view-toggle,
body.is-desktop-portrait .audio-toggle,
body.is-desktop-portrait .compact-toggle {{
  font-size: 0.72rem;
  padding: 4px 8px;
  width: 82px;
  line-height: 1.0;
}}
body.is-desktop-portrait .toolbar {{
  gap: 6px;
}}
body.is-desktop-portrait .compact-race-time {{
  font-size: 1.44rem;
}}
body.is-desktop-portrait .group-table th, body.is-desktop-portrait .group-table td,
body.is-desktop-portrait .compact-table th, body.is-desktop-portrait .compact-table td,
body.is-desktop-portrait th, body.is-desktop-portrait td {{
  font-size: inherit;
}}
body.is-desktop-portrait .compact-table th:nth-child(1), body.is-desktop-portrait .compact-table td:nth-child(1) {{
  width: 120px;
}}
body.is-desktop-portrait .compact-table th:nth-child(3), body.is-desktop-portrait .compact-table td:nth-child(3) {{
  width: 150px;
}}
body.is-desktop-portrait .compact-table th:nth-child(4), body.is-desktop-portrait .compact-table td:nth-child(4) {{
  width: 120px;
}}
body.is-desktop-portrait .compact-table-separate th:nth-child(1), body.is-desktop-portrait .compact-table-separate td:nth-child(1) {{
  width: 72px;
}}
body.is-desktop-portrait .compact-table-separate th:nth-child(3), body.is-desktop-portrait .compact-table-separate td:nth-child(3) {{
  width: 128px;
}}
body.is-desktop-portrait .compact-table-separate th:nth-child(4), body.is-desktop-portrait .compact-table-separate td:nth-child(4) {{
  width: 96px;
}}
body.is-desktop-portrait .compact-table-separate th:nth-child(5), body.is-desktop-portrait .compact-table-separate td:nth-child(5) {{
  width: 120px;
}}
body.is-desktop-portrait .compact-table-separate .compact-group-cell {{
  font-size: 2.35rem;
}}
body.is-desktop-portrait .compact-group-cell {{
  font-size: 2.8rem;
  line-height: 1.0;
}}
.tt-wrap {{
  display: flex;
  flex-direction: column;
  gap: 10px;
  flex: 1;
  min-height: 0;
}}
.tt-pane {{
  background: var(--panel);
  border: 1px solid var(--border);
  border-radius: 12px;
  padding: 8px;
  min-height: 0;
  display: flex;
  flex-direction: column;
}}
.tt-pane-top {{
  flex: 0 0 auto;
}}
.tt-pane-bottom {{
  flex: 1 1 0;
  min-height: 0;
}}
.tt-pane-title {{
  font-size: 1rem;
  font-weight: 700;
  margin-bottom: 6px;
}}
.tt-table {{
  width: 100%;
  border-collapse: collapse;
  table-layout: fixed;
  font-size: 1rem;
}}
.tt-table col.tt-col-bib {{ width: 96px; min-width: 96px; }}
.tt-table col.tt-col-start {{ width: 8ch; }}
.tt-table col.tt-col-early {{ width: 9ch; }}
.tt-table col.tt-col-stop {{ width: 9ch; }}
.tt-table col.tt-col-elapsed {{ width: 9ch; }}
.tt-table col.tt-col-name {{ width: 24ch; }}
.tt-table th, .tt-table td {{
  text-align: left;
  padding: 6px 8px;
  border-bottom: 1px solid var(--border);
  white-space: nowrap;
  overflow: hidden;
  text-overflow: ellipsis;
}}
.tt-table th {{
  background: var(--panel-2);
  position: sticky;
  top: 0;
}}
.tt-table th.tt-sortable {{
  cursor: pointer;
  user-select: none;
  -webkit-user-select: none;
}}
.tt-table th.tt-sortable:hover {{
  background: rgba(214, 228, 241, 0.98);
}}
.tt-table th.tt-num, .tt-table td.tt-num {{
  text-align: right;
  font-variant-numeric: tabular-nums;
}}
.tt-table th.tt-bib, .tt-table td.tt-bib {{
  min-width: 96px;
  width: 96px;
  overflow: visible;
  text-overflow: clip;
}}
#ttActiveRows .tt-table td.tt-bib {{
  font-size: 2rem;
  font-weight: 800;
}}
#ttCompletedRows .tt-table td.tt-bib {{
  font-weight: 800;
}}
.tt-pane-bottom .tt-table tbody tr:nth-child(odd) {{
  background: rgba(255,255,255,.98);
}}
.tt-pane-bottom .tt-table tbody tr:nth-child(even) {{
  background: rgba(243,246,249,.98);
}}
.tt-scroll {{
  overflow: auto;
  min-height: 0;
  flex: 1;
}}
.tt-pane-top .tt-scroll {{
  overflow-y: visible;
  overflow-x: auto;
  flex: 0 0 auto;
}}
.tt-pane-bottom .tt-scroll {{
  overflow-y: auto;
  overflow-x: auto;
  flex: 1 1 auto;
  min-height: 0;
}}
@media (orientation: landscape) {{
  .tt-wrap {{
    flex-direction: row;
    align-items: stretch;
  }}
  .tt-pane-top,
  .tt-pane-bottom {{
    flex: 1 1 0;
    min-width: 0;
  }}
  .tt-pane-top .tt-scroll {{
    overflow-y: auto;
    flex: 1 1 auto;
    min-height: 0;
  }}
}}
</style>
</head>
<body>
<main>
  <div class="title">
    <div class="title-left">
      <div class="compact-race-time" id="compactRaceTime">00:00:00</div>
      <div class="ip-healthbar" id="ipHealthBar"></div>
    </div>
    <div class="title-center">
      <h1 id="event" style="margin:0;font-size:1.1rem;line-height:1.1;">lapsrv</h1>
    </div>
    <div class="toolbar">
      <div class="tt-indicator" id="ttIndicator" style="display:none;">TT</div>
      <button class="warning-toggle" id="warningToggle" type="button" style="display:none;">Warning On</button>
      <button class="compact-toggle" id="compactToggle" type="button">Separate Off</button>
      <button class="view-toggle" id="viewToggle" type="button">Compact View</button>
      <button class="audio-toggle" id="audioToggle" type="button">Tone Off</button>
    </div>
  </div>
  <div class="compact-view compact-wrap">
    <div class="group-rows recent-group-rows" id="compactRecentRows"></div>
  </div>
  <div class="tt-view tt-wrap">
    <div class="tt-pane tt-pane-top">
      <div class="tt-pane-title">Approaching</div>
      <div class="tt-scroll"><div id="ttActiveRows"></div></div>
    </div>
    <div class="tt-pane tt-pane-bottom">
      <div class="tt-pane-title">Finished</div>
      <div class="tt-scroll"><div id="ttCompletedRows"></div></div>
    </div>
  </div>
  <div class="default-view group-wrap">
    <div class="group-title">Recent Groups</div>
    <div class="group-rows recent-group-rows" id="recentGroupRows"></div>
  </div>
  <div class="default-view group-wrap">
    <div class="group-title">Past Groups</div>
    <div class="group-rows past-group-rows" id="pastGroupRows"></div>
  </div>
  <div class="default-view table-wrap" id="tableWrap">
    <table class="small-table">
      <thead>
        <tr>
          <th>Bib</th>
          <th>Note</th>
          <th>Time</th>
          <th>Gap</th>
          <th>Lap</th>
          <th>Name</th>
          <th>Category</th>
          <th>Team</th>
          <th>ETA</th>
        </tr>
      </thead>
      <tbody id="rows"></tbody>
    </table>
  </div>
</main>
<script>
const selectedCategory = {category_js};
const apiPath = {json.dumps(api_path)};
let autoScroll = true;
let compactAutoScroll = true;
let viewMode = localStorage.getItem('lapsrv_view_mode') || 'full';
let compactSeparate = localStorage.getItem('lapsrv_compact_separate') === '1';
let lastIsTimeTrial = false;
let ttCompletedSort = {{ key: null, dir: 'desc' }};
let audioCtx = null;
let lastBellToneKey = '';
let audioEnabled = false;
function esc(s) {{
  return (s ?? '').toString()
    .replaceAll('&', '&amp;')
    .replaceAll('<', '&lt;')
    .replaceAll('>', '&gt;')
    .replaceAll('"', '&quot;');
}}
function shortCategory(name) {{
  return (name || '')
    .replaceAll(' (Men)', '')
    .replaceAll(' (Open)', '')
    .replaceAll(' (Women)', '-W');
}}
function updateCompactToggle() {{
  const button = document.getElementById('compactToggle');
  if (!button) return;
  button.textContent = compactSeparate ? 'Separate On' : 'Separate Off';
  button.classList.toggle('active', compactSeparate);
  button.style.display = lastIsTimeTrial && viewMode === 'tt' ? 'none' : '';
}}
function compactPos(note) {{
  const raw = (note || '').trim();
  if (!raw) return '';
  return raw.includes(':') ? raw.split(':', 1)[0] : raw;
}}
function parseClockText(text) {{
  const value = (text || '').trim();
  if (!value) return null;
  const sign = value.startsWith('-') ? -1 : 1;
  const clean = sign < 0 ? value.slice(1) : value;
  const parts = clean.split(':');
  if (parts.length < 2 || parts.length > 3) return null;
  let hours = 0, minutes = 0, seconds = 0;
  if (parts.length === 3) {{
    hours = Number(parts[0]);
    minutes = Number(parts[1]);
    seconds = Number(parts[2]);
  }} else {{
    minutes = Number(parts[0]);
    seconds = Number(parts[1]);
  }}
  if ([hours, minutes, seconds].some(Number.isNaN)) return null;
  return sign * (hours * 3600 + minutes * 60 + seconds);
}}
function compareTTValues(a, b, key) {{
  if (key === 'bib') return Number(a[key] || 0) - Number(b[key] || 0);
  if (['race_time', 'early_time', 'start_time', 'stop_time', 'elapsed'].includes(key)) {{
    const av = parseClockText(a[key]);
    const bv = parseClockText(b[key]);
    return (av ?? -Infinity) - (bv ?? -Infinity);
  }}
  return (a[key] || '').localeCompare(b[key] || '', undefined, {{ sensitivity: 'base' }});
}}
function sortTTRows(rows) {{
  const out = (rows || []).slice();
  if (!ttCompletedSort.key) return out;
  out.sort((a, b) => {{
    const cmp = compareTTValues(a, b, ttCompletedSort.key);
    return ttCompletedSort.dir === 'asc' ? cmp : -cmp;
  }});
  return out;
}}
function ttHeaderLabel(label, key, sortable) {{
  if (!sortable) return esc(label);
  if (ttCompletedSort.key !== key) return esc(label);
  return esc(label) + ' ' + (ttCompletedSort.dir === 'asc' ? '&#9650;' : '&#9660;');
}}
function trimClockFraction(text) {{
  return ((text || '').split('.', 1)[0] || text || '');
}}
function topGroupBellKey(groupTable) {{
  const headers = groupTable.headers || [];
  const rows = groupTable.rows || [];
  const headerStatus = groupTable.header_status || {{}};
  let inTopGroup = false;
  const bellHeaders = [];
  for (const row of rows) {{
    if (row.elapsed_text) {{
      if (inTopGroup) break;
      inTopGroup = true;
    }}
    if (!inTopGroup) continue;
    for (const h of headers) {{
      const cell = (row.cells || {{}})[h];
      if (!cell) continue;
      const status = headerStatus[h] || {{}};
      if (status.is_bell || status.text === '1') bellHeaders.push(h);
    }}
  }}
  bellHeaders.sort();
  return bellHeaders.join('|');
}}
function ensureAudio() {{
  if (audioCtx) return audioCtx;
  const Ctor = window.AudioContext || window.webkitAudioContext;
  if (!Ctor) return null;
  audioCtx = new Ctor();
  return audioCtx;
}}
async function armAudio() {{
  audioEnabled = true;
  updateAudioButton();
  const ctx = ensureAudio();
  if (ctx && ctx.state === 'suspended') await ctx.resume().catch(() => {{}});
}}
function playBellTone() {{
  if (!audioEnabled) return;
  const ctx = ensureAudio();
  if (!ctx) return;
  if (ctx.state === 'suspended') {{
    ctx.resume().catch(() => {{}});
  }}
  const now = ctx.currentTime;
  const freqs = [880, 1174];
  freqs.forEach((freq, i) => {{
    const osc = ctx.createOscillator();
    const gain = ctx.createGain();
    osc.type = 'sine';
    osc.frequency.setValueAtTime(freq, now);
    gain.gain.setValueAtTime(0.0001, now);
    gain.gain.exponentialRampToValueAtTime(0.08, now + 0.02 + i * 0.18);
    gain.gain.exponentialRampToValueAtTime(0.0001, now + 0.16 + i * 0.18);
    osc.connect(gain);
    gain.connect(ctx.destination);
    osc.start(now + i * 0.18);
    osc.stop(now + 0.18 + i * 0.18);
  }});
}}
function maybePlayBellTone(groupTable) {{
  const bellKey = topGroupBellKey(groupTable);
  if (bellKey && bellKey !== lastBellToneKey) playBellTone();
  lastBellToneKey = bellKey;
}}

function compactGroupTableHtml(groupTable) {{
  if (compactSeparate) return compactGroupTableHtmlSeparate(groupTable);
  const headers = groupTable.headers || [];
  const rows = groupTable.rows || [];
  const headerStatus = groupTable.header_status || {{}};
  const merged = new Map();

  rows.forEach(row => {{
    headers.forEach(h => {{
      const cell = (row.cells || {{}})[h];
      if (!cell) return;
      const status = headerStatus[h] || {{}};
      const deficitCounts = cell.deficit_counts || {{}};
      const deficitKeys = Object.keys(deficitCounts).map(v => Number(v)).sort((a, b) => a - b);
      if (deficitKeys.length) {{
        deficitKeys.forEach(deficit => {{
          const countValue = Number(deficitCounts[deficit] || 0);
          if (!countValue) return;
          const lapText = deficit > 0 ? `(-${{deficit}})` : (status.is_bell ? '1' : (status.text || ''));
          const key = `${{h}}|${{lapText}}`;
          const posText = compactPos(cell.note || cell.text || '');
          const existing = merged.get(key) || {{
            lapText,
            categoryText: shortCategory(h),
            posText,
            count: 0,
            isBell: false,
            isLead: false,
            isLapped: deficit > 0,
          }};
          existing.count += countValue;
          if (!existing.posText || (posText && posText.length < existing.posText.length)) existing.posText = posText;
          existing.isBell = existing.isBell || (deficit === 0 && status.is_bell);
          existing.isLead = existing.isLead || (deficit === 0 && cell.is_lead);
          existing.isLapped = existing.isLapped || deficit > 0;
          merged.set(key, existing);
        }});
      }} else {{
        const lapText = status.is_bell ? '1' : (status.text || '');
        const key = `${{h}}|${{lapText}}`;
        const posText = compactPos(cell.note || cell.text || '');
        const existing = merged.get(key) || {{
          lapText,
          categoryText: shortCategory(h),
          posText,
          count: 0,
          isBell: false,
          isLead: false,
          isLapped: false,
        }};
        existing.count += Number(cell.count || 0);
        if (!existing.posText || (posText && posText.length < existing.posText.length)) existing.posText = posText;
        existing.isBell = existing.isBell || status.is_bell;
        existing.isLead = existing.isLead || cell.is_lead;
        merged.set(key, existing);
      }}
    }});
  }});

  const lines = [];
  Array.from(merged.values()).forEach(entry => {{
    const rowClasses = [];
    if (entry.isBell) rowClasses.push('compact-group-bell');
    if (entry.isLead) rowClasses.push('compact-group-lead');
    if (entry.isLapped) rowClasses.push('compact-group-lapped');
    const rowClass = rowClasses.join(' ');
    lines.push(`<tr class="${{rowClass}}"><td>${{esc(entry.lapText)}}</td><td class="compact-group-cell">${{esc(entry.categoryText)}}</td><td>${{esc(entry.posText)}}</td><td>${{esc(String(entry.count))}}</td></tr>`);
  }});
  return `<table class="compact-table"><thead><tr><th>Lap</th><th>Category</th><th>Note</th><th>Count</th></tr></thead><tbody>${{lines.join('')}}</tbody></table>`;
}}
function compactGroupTableHtmlSeparate(groupTable) {{
  const headers = groupTable.headers || [];
  const rows = (groupTable.rows || []).slice().reverse();
  const headerStatus = groupTable.header_status || {{}};
  const lines = [];
  rows.forEach((row, rowIndex) => {{
    headers.forEach((h) => {{
      const cell = (row.cells || {{}})[h];
      if (!cell) return;
      const status = headerStatus[h] || {{}};
      const deficitCounts = cell.deficit_counts || {{}};
      const deficitKeys = Object.keys(deficitCounts).map(v => Number(v)).sort((a, b) => a - b);
      const pushLine = (lapText, countValue, isLapped) => {{
        if (!countValue) return;
        const rowClasses = [];
        if (!isLapped && status.is_bell) rowClasses.push('compact-group-bell');
        if (!isLapped && cell.is_lead) rowClasses.push('compact-group-lead');
        if (isLapped) rowClasses.push('compact-group-lapped');
        const rowClass = rowClasses.join(' ');
        lines.push(`<tr class="${{rowClass}}"><td>${{esc(lapText)}}</td><td class="compact-group-cell">${{esc(shortCategory(h))}}</td><td>${{esc(compactPos(cell.note || cell.text || ''))}}</td><td>${{esc(String(countValue))}}</td><td>${{esc(row.elapsed_text || '')}}</td></tr>`);
      }};
      if (deficitKeys.length) {{
        deficitKeys.forEach((deficit) => {{
          const countValue = Number(deficitCounts[deficit] || 0);
          const lapText = deficit > 0 ? `(-${{deficit}})` : (status.is_bell ? '1' : (status.text || ''));
          pushLine(lapText, countValue, deficit > 0);
        }});
      }} else {{
        const lapText = status.is_bell ? '1' : (status.text || '');
        pushLine(lapText, Number(cell.count || 0), false);
      }}
    }});
    if (rowIndex !== rows.length - 1) {{
      lines.push('<tr class="compact-separator"><td colspan="5"></td></tr>');
    }}
  }});
  return `<table class="compact-table compact-table-separate"><thead><tr><th>Lap</th><th>Category</th><th>Note</th><th>Count</th><th></th></tr></thead><tbody>${{lines.join('')}}</tbody></table>`;
}}
function ttTableHtml(rows, stopLabel, sortable=false) {{
  const viewRows = sortable ? sortTTRows(rows) : (rows || []);
  const body = viewRows.map((row) => `<tr><td class="tt-num tt-bib">${{esc(row.bib || '')}}</td><td class="tt-num">${{esc(row.start_time || '')}}</td><td class="tt-num">${{esc(row.early_time || row.race_time || '')}}</td><td class="tt-num">${{esc(row.stop_time || '')}}</td><td class="tt-num">${{esc(row.elapsed || '')}}</td><td>${{esc(((row.last_name || '') + ',' + (row.first_name || '')).replace(/^,|,$/g, ''))}}</td></tr>`).join('');
  const th = (label, key, cls='') => '<th class="' + cls + (sortable ? ' tt-sortable' : '') + '"' + (sortable ? ' data-tt-sort="' + key + '"' : '') + '>' + ttHeaderLabel(label, key, sortable) + '</th>';
  return '<table class="tt-table"><colgroup><col class="tt-col-bib"><col class="tt-col-start"><col class="tt-col-early"><col class="tt-col-stop"><col class="tt-col-elapsed"><col class="tt-col-name"></colgroup><thead><tr>' + th('BIB', 'bib', 'tt-num tt-bib') + th('Start', 'start_time', 'tt-num') + th('Early', 'early_time', 'tt-num') + th('Finish', 'stop_time', 'tt-num') + th('Elapsed', 'elapsed', 'tt-num') + th('Name', 'last_name') + '</tr></thead><tbody>' + body + '</tbody></table>';
}}
function applyViewMode() {{
  document.body.classList.toggle('compact-mode', viewMode === 'compact');
  document.body.classList.toggle('tt-mode', viewMode === 'tt');
  const button = document.getElementById('viewToggle');
  if (!button) return;
  if (lastIsTimeTrial) {{
    const nextLabel = viewMode === 'full' ? 'Compact View' : (viewMode === 'compact' ? 'TT View' : 'Full View');
    button.textContent = nextLabel;
  }} else {{
    button.textContent = viewMode === 'compact' ? 'Full View' : 'Compact View';
  }}
  updateCompactToggle();
}}
function updateAudioButton() {{
  const button = document.getElementById('audioToggle');
  if (!button) return;
  button.textContent = audioEnabled ? 'Tone On' : 'Tone Off';
  button.classList.toggle('active', audioEnabled);
}}
function updateWarningButton(enabled, visible) {{
  if (!warningToggle) return;
  warningToggle.style.display = visible ? '' : 'none';
  warningToggle.textContent = enabled ? 'Warning On' : 'Warning Off';
  warningToggle.classList.toggle('active', enabled);
}}
function updateTTIndicator(isTimeTrial) {{
  const indicator = document.getElementById('ttIndicator');
  if (!indicator) return;
  indicator.style.display = isTimeTrial ? '' : 'none';
}}
async function setWarningEnabled(enabled) {{
  await fetch('/api/warning', {{
    method: 'POST',
    headers: {{ 'Content-Type': 'application/json' }},
    body: JSON.stringify({{ enabled }}),
  }});
}}
function applyDeviceClass() {{
  const ua = navigator.userAgent || '';
  const uaDataMobile = navigator.userAgentData && typeof navigator.userAgentData.mobile === 'boolean'
    ? navigator.userAgentData.mobile
    : null;
  const mobileByUA = /iPhone|Android.+Mobile|Windows Phone|iPod/i.test(ua);
  const mobileByTouch = navigator.maxTouchPoints > 1 && Math.min(window.screen.width, window.screen.height) <= 480;
  const isMobile = uaDataMobile !== null ? uaDataMobile : (mobileByUA || mobileByTouch);
  const isDesktopPortrait = !isMobile && window.innerHeight > window.innerWidth;
  document.body.classList.toggle('is-mobile', isMobile);
  document.body.classList.toggle('is-desktop-portrait', isDesktopPortrait);
}}
function groupTableHtml(groupTable) {{
  const headers = groupTable.headers || [];
  const rows = groupTable.rows || [];
  const headerStatus = groupTable.header_status || {{}};
  const ths = ['<th>HH:MM:SS</th>', '<th>Gap</th>', '<th>Total</th>']
    .concat(headers.map(h => `<th>${{esc(h)}}</th>`))
    .join('');
  const statusTds = ['<th></th>', '<th></th>', '<th></th>']
    .concat(headers.map(h => {{
      const status = headerStatus[h] || {{}};
      const cls = status.is_bell ? 'group-header-status bell' : 'group-header-status';
      return `<th class="${{cls}}">${{esc(status.text || '')}}</th>`;
    }}))
    .join('');
  const trs = rows.map(row => {{
    const tds = [`<td>${{esc(row.elapsed_text || '')}}</td>`, `<td>${{esc(row.group_gap || '')}}</td>`, `<td>${{esc(row.total ?? '')}}</td>`]
      .concat(headers.map(h => {{
        const cell = (row.cells || {{}})[h];
        if (!cell) return '<td class="group-empty"></td>';
        const cls = ['group-cell'];
        if (cell.is_lead) cls.push('lead');
        if (cell.is_lapped) cls.push('lapped');
        const note = esc(cell.note || cell.text || '');
        const count = (cell.count || 0) > 1 ? esc(String(cell.count)) : '';
        return `<td class="${{cls.join(' ')}}"><div class="group-cell-content"><span class="group-cell-note">${{note}}</span><span class="group-cell-count">${{count}}</span></div></td>`;
      }}))
      .join('');
    return `<tr>${{tds}}</tr>`;
  }}).join('');
  return `<table class="group-table"><thead><tr>${{statusTds}}</tr><tr>${{ths}}</tr></thead><tbody>${{trs}}</tbody></table>`;
}}
function statusHtml(s) {{
  const endpoint = s.endpoint ? ` ${{esc(s.endpoint)}}` : '';
  return `<span class="badge state-${{esc(s.state)}}">${{esc(s.name)}} [${{esc(s.state)}}]${{endpoint}}</span>`;
}}
function ipHealthHtml(items) {{
  return (items || []).map(item => `<span class="ip-pill ${{esc(item.state)}}">${{esc(item.address)}}</span>`).join('');
}}
function rowHtml(r) {{
  return `<tr>
    <td>${{esc(r.bib)}}</td>
    <td class="note">${{esc(r.note)}}</td>
    <td>${{esc(r.time)}}</td>
    <td>${{esc(r.gap)}}</td>
    <td>${{esc(r.lap ?? '')}}</td>
    <td>${{esc(r.name)}}</td>
    <td>${{esc(r.category)}}</td>
    <td>${{esc(r.team)}}</td>
    <td>${{esc(r.eta)}}</td>
  </tr>`;
}}
const tableWrap = document.getElementById('tableWrap');
const recentGroupRowsWrap = document.getElementById('recentGroupRows');
const pastGroupRowsWrap = document.getElementById('pastGroupRows');
const compactRecentRowsWrap = document.getElementById('compactRecentRows');
const ttCompletedRowsEl = document.getElementById('ttCompletedRows');
const compactToggle = document.getElementById('compactToggle');
const warningToggle = document.getElementById('warningToggle');
const viewToggle = document.getElementById('viewToggle');
const audioToggle = document.getElementById('audioToggle');
const ipHealthBar = document.getElementById('ipHealthBar');
applyDeviceClass();
applyViewMode();
updateAudioButton();
updateCompactToggle();
window.addEventListener('resize', applyDeviceClass);
if (ttCompletedRowsEl) {{
  ttCompletedRowsEl.addEventListener('click', (event) => {{
    const th = event.target.closest('th[data-tt-sort]');
    if (!th) return;
    const key = th.dataset.ttSort;
    if (ttCompletedSort.key === key) ttCompletedSort.dir = ttCompletedSort.dir === 'asc' ? 'desc' : 'asc';
    else {{
      ttCompletedSort.key = key;
      ttCompletedSort.dir = ['last_name', 'first_name'].includes(key) ? 'asc' : 'desc';
    }}
    ttCompletedRowsEl.innerHTML = ttTableHtml(window.__ttCompletedRows || [], 'Stop', true);
  }});
}}
if (compactToggle) {{
  compactToggle.addEventListener('click', () => {{
    compactSeparate = !compactSeparate;
    localStorage.setItem('lapsrv_compact_separate', compactSeparate ? '1' : '0');
    updateCompactToggle();
    compactRecentRowsWrap.innerHTML = compactGroupTableHtml(window.__recentGroupTable || {{ headers: [], rows: [], header_status: {{}} }});
  }});
}}
warningToggle.addEventListener('click', async () => {{
  const enabled = !(warningToggle.classList.contains('active'));
  await setWarningEnabled(enabled);
  updateWarningButton(enabled, true);
}});
viewToggle.addEventListener('click', () => {{
  if (lastIsTimeTrial) {{
    viewMode = viewMode === 'full' ? 'compact' : (viewMode === 'compact' ? 'tt' : 'full');
  }} else {{
    viewMode = viewMode === 'compact' ? 'full' : 'compact';
  }}
  localStorage.setItem('lapsrv_view_mode', viewMode);
  applyViewMode();
}});
audioToggle.addEventListener('click', async () => {{
  if (audioEnabled) {{
    audioEnabled = false;
    updateAudioButton();
    return;
  }}
  await armAudio();
  playBellTone();
}});
tableWrap.addEventListener('scroll', () => {{
  const remaining = tableWrap.scrollHeight - tableWrap.scrollTop - tableWrap.clientHeight;
  autoScroll = remaining < 40;
}});
recentGroupRowsWrap.addEventListener('scroll', () => {{
  // recent pane stays top-anchored
}});
pastGroupRowsWrap.addEventListener('scroll', () => {{
  // debug pane; no auto-follow state needed
}});
compactRecentRowsWrap.addEventListener('scroll', () => {{
  const remaining = compactRecentRowsWrap.scrollHeight - compactRecentRowsWrap.scrollTop - compactRecentRowsWrap.clientHeight;
  compactAutoScroll = remaining < 40;
}});
async function refresh() {{
  const res = await fetch(apiPath, {{cache: 'no-store'}});
  const data = await res.json();
  document.getElementById('event').textContent = data.event_name || 'lapsrv';
  const compactRaceTimeEl = document.getElementById('compactRaceTime');
  if (compactRaceTimeEl) compactRaceTimeEl.textContent = trimClockFraction(data.cur_race_time || data.current_race_time || data.race_clock || '00:00:00');
  lastIsTimeTrial = !!data.is_time_trial;
  updateTTIndicator(lastIsTimeTrial);
  if (!lastIsTimeTrial && viewMode === 'tt') {{
    viewMode = 'full';
    localStorage.setItem('lapsrv_view_mode', viewMode);
  }}
  applyViewMode();
  ipHealthBar.innerHTML = ipHealthHtml(data.ip_health || []);
  updateWarningButton(!!data.warning_enabled, (data.ip_health || []).length > 0);
  const subsystems = data.subsystems || {{}};
  const recentGroupTable = data.recent_group_table || {{headers: [], header_status: {{}}, rows: []}};
  window.__recentGroupTable = recentGroupTable;
  recentGroupRowsWrap.innerHTML = groupTableHtml(recentGroupTable);
  compactRecentRowsWrap.innerHTML = compactGroupTableHtml(recentGroupTable);
  if (compactAutoScroll) {{
    compactRecentRowsWrap.scrollTop = compactRecentRowsWrap.scrollHeight;
  }}
  window.__ttCompletedRows = data.tt_completed_rows || [];
  document.getElementById('ttActiveRows').innerHTML = ttTableHtml(data.tt_active_rows || [], 'Early');
  document.getElementById('ttCompletedRows').innerHTML = ttTableHtml(window.__ttCompletedRows, 'Stop', true);
  maybePlayBellTone(recentGroupTable);
  pastGroupRowsWrap.innerHTML = groupTableHtml(data.past_group_table || {{headers: [], header_status: {{}}, rows: []}});
  document.getElementById('rows').innerHTML = (data.passings || []).map(rowHtml).join('');
  if (autoScroll) {{
    tableWrap.scrollTop = tableWrap.scrollHeight;
  }}
}}
refresh();
setInterval(refresh, 1000);
</script>
</body>
</html>"""


class TkMonitor:
    def __init__(self, state: SharedState, on_close: Optional[callable] = None) -> None:
        self.state = state
        self.on_close = on_close
        self.thread: Optional[threading.Thread] = None
        self.logger = logging.getLogger("lapsrv.tk")
        self._closed = False
        self._after_id: Optional[str] = None

    def start(self) -> None:
        if not os.environ.get("DISPLAY") and os.name != "nt":
            self.logger.warning("DISPLAY is not set, skipping tkinter UI")
            return
        self.thread = threading.Thread(target=self._run, name="lapsrv-tk", daemon=True)
        self.thread.start()

    def _run(self) -> None:
        root = tk.Tk()
        root.title("lapsrv")
        root.geometry("1400x900")

        def handle_close() -> None:
            if self._closed:
                return
            self._closed = True
            if self._after_id is not None:
                try:
                    root.after_cancel(self._after_id)
                except Exception:
                    pass
                self._after_id = None
            if self.on_close:
                try:
                    self.on_close()
                except Exception:
                    self.logger.exception("tk close callback failed")
            try:
                root.quit()
            except Exception:
                pass
            try:
                root.destroy()
            except Exception:
                pass

        root.protocol("WM_DELETE_WINDOW", handle_close)

        header = ttk.Label(root, text="lapsrv", font=("TkDefaultFont", 18, "bold"))
        header.pack(fill="x", padx=8, pady=(8, 4))

        ip_frame = tk.Frame(root)
        ip_frame.pack(fill="x", padx=8, pady=(0, 4))

        recent_frame = ttk.LabelFrame(root, text="Recent Groups")
        recent_frame.pack(fill="both", expand=True, padx=8, pady=(8, 4))
        recent_group_text = tk.Text(recent_frame, wrap="none", height=15, font=("Courier New", 15), state="disabled")
        recent_group_text.tag_configure("header", font=("Courier New", 15, "bold"))
        recent_group_text.tag_configure("lead", background="#facc15", foreground="#111827")
        recent_group_text.tag_configure("bell", background="#facc15", foreground="#111827")
        recent_group_text.tag_configure("lapped", background="#4b5563", foreground="#e5edf6")
        recent_group_text.tag_configure("normal", foreground="#111827")
        recent_group_scroll = ttk.Scrollbar(recent_frame, orient="vertical", command=recent_group_text.yview)
        recent_group_text.configure(yscrollcommand=recent_group_scroll.set)
        recent_group_text.pack(side="left", fill="both", expand=True)
        recent_group_scroll.pack(side="right", fill="y")

        past_frame = ttk.LabelFrame(root, text="Past Groups")
        past_frame.pack(fill="both", expand=True, padx=8, pady=(0, 4))
        past_group_text = tk.Text(past_frame, wrap="none", height=13, font=("Courier New", 13), state="disabled")
        past_group_text.tag_configure("header", font=("Courier New", 13, "bold"))
        past_group_text.tag_configure("lead", background="#facc15", foreground="#111827")
        past_group_text.tag_configure("bell", background="#facc15", foreground="#111827")
        past_group_text.tag_configure("lapped", background="#4b5563", foreground="#e5edf6")
        past_group_text.tag_configure("normal", foreground="#111827")
        past_group_scroll = ttk.Scrollbar(past_frame, orient="vertical", command=past_group_text.yview)
        past_group_text.configure(yscrollcommand=past_group_scroll.set)
        past_group_text.pack(side="left", fill="both", expand=True)
        past_group_scroll.pack(side="right", fill="y")

        pass_frame = ttk.LabelFrame(root, text="Recent Passings")
        pass_frame.pack(fill="both", expand=False, padx=8, pady=(0, 8))
        cols = ("Bib", "Note", "Time", "Gap", "Lap", "Name", "Category", "Team", "ETA")
        pass_tree = ttk.Treeview(pass_frame, columns=cols, show="headings", height=6)
        for col in cols:
            pass_tree.heading(col, text=col)
            width = 90
            if col in {"Name", "Team", "Category"}:
                width = 180
            pass_tree.column(col, width=width, anchor="w" if col in {"Name", "Team", "Category"} else "center")
        scrollbar = ttk.Scrollbar(pass_frame, orient="vertical", command=pass_tree.yview)
        pass_tree.configure(yscrollcommand=scrollbar.set)
        pass_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        def repaint() -> None:
            if self._closed or not root.winfo_exists():
                return
            snap = self.state.snapshot()
            subsystems = snap.get("subsystems", {})
            status_parts = []
            for key in ("rfid", "announcer", "lapcounter"):
                sub = subsystems.get(key)
                if not sub:
                    continue
                endpoint = f" {sub['endpoint']}" if sub.get("endpoint") else ""
                status_parts.append(f"{sub['name']} [{sub['state']}]{{{endpoint.strip()}}}" if endpoint else f"{sub['name']} [{sub['state']}]")
            status_text = "   ".join(status_parts)
            header.config(text=f"{snap['event_name'] or 'lapsrv'}    Race {snap['current_race_time'] or ''}    {status_text}")
            for child in ip_frame.winfo_children():
                child.destroy()
            for item in snap.get("ip_health", []):
                bg = "#86efac" if item.get("state") == "ok" else ("#fde68a" if item.get("state") == "missed" else "#fbcfe8")
                fg = "#111827"
                lbl = tk.Label(ip_frame, text=item.get("address", ""), bg=bg, fg=fg, padx=10, pady=4, relief="groove", borderwidth=1)
                lbl.pack(side="left", padx=(0, 6))

            def render_group_text(widget: tk.Text, group_table: dict[str, Any]) -> None:
                headers = group_table.get("headers", [])
                header_status = group_table.get("header_status", {})
                rows = group_table.get("rows", [])
                widths = {"HH:MM:SS": 9, "Gap": 8, "Total": 5}
                for header_name in headers:
                    status_text = (header_status.get(header_name) or {}).get("text", "")
                    max_len = max([len(header_name), len(status_text)] + [len((row.get("cells", {}).get(header_name) or {}).get("text", "")) for row in rows])
                    widths[header_name] = max(12, min(max_len + 2, 28))
                widths["Total"] = max(widths["Total"], max([len("Total")] + [len(str(row.get("total", ""))) for row in rows] if rows else [len("Total")]))
                group_lines = []
                header_line = f'{"HH:MM:SS":<{widths["HH:MM:SS"]}}  {"Gap":<{widths["Gap"]}}  {"Total":<{widths["Total"]}}'
                for header_name in headers:
                    header_line += f'  {header_name:<{widths[header_name]}}'
                status_line = f'{"":<{widths["HH:MM:SS"]}}  {"":<{widths["Gap"]}}  {"":<{widths["Total"]}}'
                bell_ranges = []
                offset = widths["HH:MM:SS"] + 2 + widths["Gap"] + 2 + widths["Total"]
                for header_name in headers:
                    offset += 2
                    status = header_status.get(header_name) or {}
                    status_text = f'{status.get("text", ""):<{widths[header_name]}}'
                    status_line += f'  {status_text}'
                    if status.get("is_bell"):
                        bell_ranges.append((offset, offset + len(status_text)))
                    offset += len(status_text)
                group_lines.append((status_line, [], bell_ranges, []))
                group_lines.append((header_line, [], [], []))
                for row in rows:
                    text_parts = [
                        f'{row.get("elapsed_text", ""):<{widths["HH:MM:SS"]}}',
                        f'{row.get("group_gap", ""):<{widths["Gap"]}}',
                        f'{row.get("total", ""):<{widths["Total"]}}',
                    ]
                    lead_ranges = []
                    lapped_ranges = []
                    for header_name in headers:
                        cell = (row.get("cells", {}) or {}).get(header_name)
                        if cell:
                            note_text = cell.get("note", cell.get("text", ""))
                            count_text = str(cell.get("count", "")) if cell.get("count", 0) > 1 else ""
                            if cell.get("is_lapped"):
                                cell_text = f'{note_text} {count_text}'.rstrip()
                                padded = f'{cell_text:>{widths[header_name]}}'
                            elif count_text:
                                gap_width = max(widths[header_name] - len(note_text) - len(count_text), 1)
                                padded = f'{note_text}{" " * gap_width}{count_text}'
                            else:
                                padded = f'{note_text:<{widths[header_name]}}'
                        else:
                            padded = f'{"":<{widths[header_name]}}'
                        start = sum(len(part) for part in text_parts) + (2 * len(text_parts))
                        text_parts.append(padded)
                        if cell and cell.get("is_lead"):
                            lead_ranges.append((start, start + len(padded)))
                        elif cell and cell.get("is_lapped"):
                            lapped_ranges.append((start, start + len(padded)))
                    group_lines.append(("  ".join(text_parts), lead_ranges, [], lapped_ranges))
                widget.configure(state="normal")
                widget.delete("1.0", "end")
                widget.tag_remove("lead", "1.0", "end")
                widget.tag_remove("bell", "1.0", "end")
                widget.tag_remove("lapped", "1.0", "end")
                for line, lead_ranges, bell_ranges, lapped_ranges in group_lines:
                    start_idx = widget.index("end-1c")
                    widget.insert("end", line + "\n")
                    line_no = start_idx.split('.')[0]
                    for col_start, col_end in lead_ranges:
                        if col_end > col_start:
                            widget.tag_add("lead", f"{line_no}.{col_start}", f"{line_no}.{col_end}")
                    for col_start, col_end in bell_ranges:
                        if col_end > col_start:
                            widget.tag_add("bell", f"{line_no}.{col_start}", f"{line_no}.{col_end}")
                    for col_start, col_end in lapped_ranges:
                        if col_end > col_start:
                            widget.tag_add("lapped", f"{line_no}.{col_start}", f"{line_no}.{col_end}")
                widget.configure(state="disabled")
                widget.see("1.0")

            render_group_text(recent_group_text, snap.get("recent_group_table", {"headers": [], "header_status": {}, "rows": []}))
            render_group_text(past_group_text, snap.get("past_group_table", {"headers": [], "header_status": {}, "rows": []}))
            for item in pass_tree.get_children():
                pass_tree.delete(item)
            for row in snap["passings"]:
                pass_tree.insert(
                    "",
                    "end",
                    values=(
                        row["bib"],
                        row["note"],
                        row["time"],
                        row["gap"],
                        "" if row["lap"] is None else row["lap"],
                        row["name"],
                        row["category"],
                        row["team"],
                        row["eta"],
                    ),
                )
            children = pass_tree.get_children()
            if children:
                pass_tree.see(children[-1])
            if not self._closed and root.winfo_exists():
                self._after_id = root.after(750, repaint)

        try:
            repaint()
            root.mainloop()
        finally:
            handle_close()


class LapsrvApp:
    def __init__(self, args: argparse.Namespace) -> None:
        self.args = args
        self.config = RaceConfig.load(Path(args.xlsx))
        self.state = SharedState(self.config, show_unknown_tags=args.show_unknown_tags)
        self.state.tt_export_path = Path(args.xlsx).with_name(Path(args.xlsx).stem + '-tt-results.csv')
        self.state.tt_xlsx_path = Path(args.xlsx).with_name(Path(args.xlsx).stem + '-tt-results.xlsx')
        self.state.road_export_path = Path(args.xlsx).with_name(Path(args.xlsx).stem + '-road-passings.csv')
        self.state.road_xlsx_path = Path(args.xlsx).with_name(Path(args.xlsx).stem + '-road-passings.xlsx')
        self.jchip = JChipServer(self.state, args.listen_host, args.port)
        self.crossmgr = CrossMgrClient(self.state, args.crossmgr_host, args.crossmgr_port)
        self.web = WebServer(self.state, args.web_host, args.web)
        self.ping = PingMonitor(self.state, list(args.ip_address or []))
        self.loop: Optional[asyncio.AbstractEventLoop] = None
        self.tk = TkMonitor(self.state, on_close=self.request_stop)
        self.stop_event = asyncio.Event()
        self.logger = logging.getLogger("lapsrv.app")
        print(f"race mode: {'tt' if self.config.time_trial else 'road'}", file=sys.stderr, flush=True)

    def request_stop(self) -> None:
        self.logger.info("shutdown requested")
        if self.loop is not None:
            self.loop.call_soon_threadsafe(self.stop_event.set)
        else:
            self.stop_event.set()

    async def _stop_with_timeout(self, name: str, awaitable: Any, timeout: float = 2.0) -> None:
        self.logger.info("shutdown: stopping %s", name)
        try:
            await asyncio.wait_for(awaitable, timeout=timeout)
            self.logger.info("shutdown: stopped %s", name)
        except asyncio.TimeoutError:
            self.logger.warning("shutdown: timed out stopping %s", name)
        except Exception:
            self.logger.exception("shutdown: error stopping %s", name)

    async def run(self) -> None:
        loop = asyncio.get_running_loop()
        self.loop = loop
        for sig in (signal.SIGINT, signal.SIGTERM):
            try:
                loop.add_signal_handler(sig, self.request_stop)
            except NotImplementedError:
                pass

        if not self.args.no_gui:
            self.tk.start()

        await self.jchip.start()
        await self.web.start()
        crossmgr_task = asyncio.create_task(self.crossmgr.run(), name="crossmgr-client")
        ping_task = asyncio.create_task(self.ping.run(), name="ping-monitor")
        try:
            await self.stop_event.wait()
        finally:
            self.logger.info("shutdown: begin")
            await self._stop_with_timeout("crossmgr", self.crossmgr.stop())
            await self._stop_with_timeout("ping", self.ping.stop())
            self.logger.info("shutdown: cancelling crossmgr task")
            crossmgr_task.cancel()
            ping_task.cancel()
            try:
                await asyncio.wait_for(asyncio.gather(crossmgr_task, ping_task, return_exceptions=True), timeout=2.0)
                self.logger.info("shutdown: background tasks cancelled")
            except asyncio.TimeoutError:
                self.logger.warning("shutdown: timed out waiting for background tasks")
            await self._stop_with_timeout("web", self.web.stop())
            await self._stop_with_timeout("jchip", self.jchip.stop())
            if self.config.time_trial:
                try:
                    exported = self.state.export_tt_xlsx()
                    if exported:
                        self.logger.info("tt results xlsx exported: %s", exported)
                except Exception:
                    self.logger.exception("failed to export tt results xlsx")
            else:
                try:
                    exported = self.state.export_road_xlsx()
                    if exported:
                        self.logger.info("road passings xlsx exported: %s", exported)
                except Exception:
                    self.logger.exception("failed to export road passings xlsx")
            self.logger.info("shutdown: complete")


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="lapsrv early-read monitor for CrossMgr races")
    parser.add_argument("--xlsx", required=True, help="CrossMgr XLSX configuration file")
    parser.add_argument("--crossmgr", default="localhost:8765", help="CrossMgr host[:port], default localhost:8765")
    parser.add_argument("--port", type=int, default=DEFAULT_JCHIP_PORT, help=f"JChip listen port, default {DEFAULT_JCHIP_PORT}")
    parser.add_argument("--web", type=int, default=DEFAULT_WEB_PORT, help=f"Web UI port, default {DEFAULT_WEB_PORT}")
    parser.add_argument("--listen-host", default="0.0.0.0", help="JChip listen host, default 0.0.0.0")
    parser.add_argument("--web-host", default="0.0.0.0", help="Web listen host, default 0.0.0.0")
    parser.add_argument("--ip_address", nargs="*", default=[], help="Up to three IP addresses to ping once per second")
    parser.add_argument("--group-gap-seconds", type=float, default=GROUP_GAP_SECONDS, help=f"Gap in seconds that starts a new group, default {GROUP_GAP_SECONDS}")
    parser.add_argument("--group-max_age-seconds", type=float, default=GROUP_MAX_AGE_SECONDS, help=f"Maximum age in seconds to retain a group, default {GROUP_MAX_AGE_SECONDS}")
    parser.add_argument("--no-gui", action="store_true", help="Disable the tkinter window")
    parser.add_argument(
        "--show_unknown_tags",
        action="store_true",
        help="Show unknown tags using the first 12 EPC characters; otherwise ignore them",
    )
    parser.add_argument("--log-level", default="INFO", help="Logging level, default INFO")
    return parser


async def async_main(args: argparse.Namespace) -> None:
    app = LapsrvApp(args)
    await app.run()


def main() -> None:
    global GROUP_GAP_SECONDS, GROUP_MAX_AGE_SECONDS
    parser = build_arg_parser()
    args = parser.parse_args()
    args.crossmgr_host, args.crossmgr_port = parse_host_port(args.crossmgr, DEFAULT_CROSSMGR_PORT)
    if len(args.ip_address) > 3:
        parser.error("--ip_address accepts at most 3 addresses")
    if args.group_gap_seconds < 0:
        parser.error("--group-gap-seconds must be >= 0")
    if args.group_max_age_seconds <= 0:
        parser.error("--group-max_age-seconds must be > 0")

    GROUP_GAP_SECONDS = float(args.group_gap_seconds)
    GROUP_MAX_AGE_SECONDS = float(args.group_max_age_seconds)

    logging.basicConfig(
        level=getattr(logging, str(args.log_level).upper(), logging.INFO),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    logging.getLogger("websockets").setLevel(logging.WARNING)
    logging.getLogger("aiohttp.access").setLevel(logging.WARNING)

    try:
        asyncio.run(async_main(args))
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
